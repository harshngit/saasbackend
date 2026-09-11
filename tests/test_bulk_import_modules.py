"""Comprehensive test suite for Bulk Excel/CSV Import and Template Download in Customer, Purchase, and Order modules."""

import io
import os
import sys
import uuid
from openpyxl import Workbook, load_workbook

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from fastapi.testclient import TestClient
from app.main import app
from app.core.database import SessionLocal
from app.models.customer import Customer
from app.models.product import Product, ProductVariant
from app.models.purchase_invoice import PurchaseInvoice
from app.models.sales_order import SalesOrder
from app.models.supplier import Supplier
from app.models.warehouse import Warehouse
from app.seed import main as seed_main

seed_main()
client = TestClient(app)

passed_count = 0
failed_count = 0


def check(description: str, condition: bool, extra: str = ""):
    global passed_count, failed_count
    if condition:
        print(f"  PASS  {description}")
        passed_count += 1
    else:
        print(f"  FAIL  {description} {extra}")
        failed_count += 1


def register_org(name_prefix: str = "Import Test Org"):
    email = f"import_admin_{uuid.uuid4().hex[:8]}@example.com"
    r = client.post(
        "/auth/register",
        json={
            "organization_name": f"{name_prefix} {uuid.uuid4().hex[:6]}",
            "admin_name": "Import Admin",
            "email": email,
            "password": "Password123!",
            "role": "admin",
        },
    )
    assert r.status_code == 201, r.text
    token = r.json()["tokens"]["access_token"]
    org_id = r.json()["user"]["organization_id"]
    headers = {"Authorization": f"Bearer {token}"}
    return org_id, headers


def create_xlsx_bytes(columns: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_csv_bytes(columns: list[str], rows: list[list]) -> bytes:
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for r in rows:
        writer.writerow(r)
    return buf.getvalue().encode("utf-8")


def run_customer_import_tests(org_id: str, headers: dict):
    print("\n--- Testing Customer Bulk Import & Template ---")

    # 1. Template download verification
    r = client.get("/customers/import/template", headers=headers)
    check("GET /customers/import/template -> 200", r.status_code == 200)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    template_cols = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    expected_cols = [
        "customer_name",
        "customer_type",
        "mobile_number",
        "email_address",
        "gstin_tax_id",
        "primary_contact_person",
        "shipping_address",
        "city",
        "payment_terms",
        "credit_limit",
    ]
    check(
        "Customer template has EXACTLY 10 expected columns in order",
        template_cols == expected_cols,
        f"got: {template_cols}",
    )

    # 2. Valid multi-row XLSX import
    valid_rows = [
        [
            "Customer One Pvt Ltd",
            "business",
            "9876543210",
            "cust1@example.com",
            "27AAAAA0000A1Z1",
            "Amit Patel",
            "123 Industrial Area",
            "Pune",
            "net_30",
            50000.0,
        ],
        [
            "Customer Two Enterprises",
            "dealer",
            "9876543211",
            "cust2@example.com",
            "27AAAAA0000A1Z2",
            "Suresh Kumar",
            "456 Market Road",
            "Mumbai",
            "due_on_receipt",
            25000.0,
        ],
    ]
    xlsx_bytes = create_xlsx_bytes(expected_cols, valid_rows)
    r = client.post(
        "/customers/import",
        headers=headers,
        files={"file": ("customers.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /customers/import valid XLSX -> 200", r.status_code == 200, r.text)
    data = r.json()
    check("Customer import 2 successes, 0 errors", data["success_count"] == 2 and data["error_count"] == 0, str(data))
    check("Returned 2 created IDs", len(data["created_ids"]) == 2)

    # Verify DB records
    db = SessionLocal()
    c1 = db.query(Customer).filter(Customer.name == "Customer One Pvt Ltd", Customer.organization_id == org_id).first()
    check("Customer 1 persisted in DB with auto CUST code", c1 is not None and c1.customer_id is not None)
    check("Customer 1 credit limit stored", c1.credit_limit == 50000.0)
    db.close()

    # 3. Valid CSV import
    csv_rows = [
        [
            "Customer Three CSV",
            "distributor",
            "9876543212",
            "cust3@example.com",
            "27AAAAA0000A1Z3",
            "Vijay Singh",
            "789 Highway Link",
            "Nagpur",
            "net_45",
            100000.0,
        ]
    ]
    csv_bytes = create_csv_bytes(expected_cols, csv_rows)
    r = client.post(
        "/customers/import",
        headers=headers,
        files={"file": ("customers.csv", csv_bytes, "text/csv")},
    )
    check("POST /customers/import valid CSV -> 200", r.status_code == 200, r.text)
    check("CSV import 1 success", r.json()["success_count"] == 1)

    # 4. Invalid rows (missing customer_name, invalid customer_type, negative credit_limit)
    bad_rows = [
        ["", "Retailer", "9999999999", "bad1@ex.com", "", "", "", "", "", 1000.0],  # missing name (row 2)
        ["Customer Bad Type", "Supermarket", "9999999998", "bad2@ex.com", "", "", "", "", "", 1000.0],  # invalid type (row 3)
        ["Customer Bad Credit", "Retailer", "9999999997", "bad3@ex.com", "", "", "", "", "", -500.0],  # negative credit (row 4)
    ]
    bad_xlsx = create_xlsx_bytes(expected_cols, bad_rows)
    r = client.post(
        "/customers/import",
        headers=headers,
        files={"file": ("bad_customers.xlsx", bad_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /customers/import bad rows -> 200 with row errors", r.status_code == 200)
    res = r.json()
    check("Recorded 3 errors, 0 successes", res["error_count"] == 3 and res["success_count"] == 0, str(res))
    err_rows = [e["row"] for e in res["errors"]]
    check("Errors accurately reference rows 2, 3, 4", err_rows == [2, 3, 4], str(err_rows))


def run_purchase_import_tests(org_id: str, headers: dict):
    print("\n--- Testing Purchase Bulk Import & Template ---")

    # Setup supplier, warehouse, products
    db = SessionLocal()
    supp1 = Supplier(
        organization_id=org_id,
        name="Apex Vendor Ltd",
        phone="9111111111",
        email="vendor1@apex.com",
        gst_number="27SUPP0000A1Z5",
        is_active=True,
    )
    wh1 = Warehouse(
        organization_id=org_id,
        name="Main Central Warehouse",
        code="WH-CENTRAL",
        is_active=True,
    )
    prod1 = Product(
        organization_id=org_id,
        name="Industrial Motor 5HP",
        sku="MOT-5HP-01",
        price=5000.0,
    )
    prod2 = Product(
        organization_id=org_id,
        name="Hydraulic Valve V2",
        sku="VAL-HYD-02",
        price=1200.0,
    )
    db.add_all([supp1, wh1, prod1, prod2])
    db.commit()
    db.refresh(supp1)
    db.refresh(wh1)
    db.refresh(prod1)
    db.refresh(prod2)

    var1 = ProductVariant(
        product_id=prod2.id,
        name="10mm Fitting",
        sku="VAL-HYD-02-10MM",
    )
    db.add(var1)
    db.commit()
    db.refresh(var1)

    supp1_id = supp1.id
    wh1_id = wh1.id
    prod1_id = prod1.id
    prod2_id = prod2.id
    var1_id = var1.id
    db.close()

    # 1. Template download verification
    r = client.get("/purchases/import/template", headers=headers)
    check("GET /purchases/import/template -> 200", r.status_code == 200)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    template_cols = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    expected_cols = [
        "invoice_number",
        "supplier_id",
        "invoice_date",
        "warehouse_id",
        "product_id",
        "variant_id",
        "quantity",
        "purchase_price",
        "item_tax_rate",
        "batch_number",
    ]
    check(
        "Purchase template has EXACTLY 10 expected columns in order",
        template_cols == expected_cols,
        f"got: {template_cols}",
    )

    # 2. Multi-item purchase grouped by invoice_number + supplier_id using human-readable codes
    rows = [
        # Purchase 1 (2 items)
        ["INV-PUR-001", "27SUPP0000A1Z5", "2026-09-11", "WH-CENTRAL", "MOT-5HP-01", "", 10, 3500.0, 18.0, "BATCH-M1"],
        ["INV-PUR-001", "27SUPP0000A1Z5", "2026-09-11", "WH-CENTRAL", "VAL-HYD-02", "VAL-HYD-02-10MM", 20, 800.0, 12.0, "BATCH-V1"],
        # Purchase 2 (1 item)
        ["INV-PUR-002", supp1_id, "2026-09-11", wh1_id, prod1_id, "", 5, 3400.0, 18.0, "BATCH-M2"],
    ]
    xlsx_bytes = create_xlsx_bytes(expected_cols, rows)
    r = client.post(
        "/purchases/import",
        headers=headers,
        files={"file": ("purchases.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /purchases/import grouped purchases -> 200", r.status_code == 200, r.text)
    data = r.json()
    check("Created 2 purchases (total 3 item rows)", data["success_count"] == 2 and data["total_rows"] == 3, str(data))

    # Verify DB purchase and items
    db = SessionLocal()
    p1 = db.query(PurchaseInvoice).filter(PurchaseInvoice.invoice_number == "INV-PUR-001", PurchaseInvoice.organization_id == org_id).first()
    check("Purchase 1 exists in DB", p1 is not None)
    check("Purchase 1 has exactly 2 items", len(p1.items) == 2)
    check("Purchase 1 total calculated correctly", p1.total == round((10 * 3500 * 1.18) + (20 * 800 * 1.12), 2))
    check("Purchase 1 in draft status with no stock added", p1.status == "draft" and p1.stock_added is False)
    db.close()

    # 3. Invalid rows (unknown supplier, unknown product, conflicting warehouse in group)
    bad_rows = [
        ["INV-BAD-01", "NON-EXISTENT-SUPP", "2026-09-11", "", "MOT-5HP-01", "", 5, 100.0, 0.0, ""],  # bad supplier
        ["INV-BAD-02", supp1_id, "2026-09-11", "", "NON-EXISTENT-PROD", "", 5, 100.0, 0.0, ""],  # bad product
        ["INV-BAD-03", supp1_id, "2026-09-11", "WH-CENTRAL", "MOT-5HP-01", "", 5, 100.0, 0.0, ""],  # item 1
        ["INV-BAD-03", supp1_id, "2026-09-11", "WH-OTHER-DIFF", "MOT-5HP-01", "", 5, 100.0, 0.0, ""],  # conflicting wh
    ]
    bad_xlsx = create_xlsx_bytes(expected_cols, bad_rows)
    r = client.post(
        "/purchases/import",
        headers=headers,
        files={"file": ("bad_purchases.xlsx", bad_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /purchases/import bad rows -> 200 with detailed errors", r.status_code == 200)
    res = r.json()
    check("Errors recorded for invalid purchases", res["error_count"] >= 3, str(res))


def run_order_import_tests(org_id: str, headers: dict):
    print("\n--- Testing Order Bulk Import & Template ---")

    # Setup customer, products
    db = SessionLocal()
    cust1 = Customer(
        organization_id=org_id,
        name="Acme Retail Stores",
        phone="9222222222",
        email="acme@retail.com",
        delivery_address="77 Commercial Street, Bengaluru",
        is_active=True,
    )
    prod1 = Product(
        organization_id=org_id,
        name="Standard Office Chair",
        sku="CHR-STD-01",
        price=2500.0,
        tax_rate=18.0,
    )
    prod2 = Product(
        organization_id=org_id,
        name="Executive Desk Pro",
        sku="DSK-PRO-02",
        price=8000.0,
        tax_rate=18.0,
    )
    db.add_all([cust1, prod1, prod2])
    db.commit()
    db.refresh(cust1)
    db.refresh(prod1)
    db.refresh(prod2)
    cust1_id = cust1.id
    prod1_id = prod1.id
    prod2_id = prod2.id
    db.close()

    # 1. Template download verification
    r = client.get("/orders/import/template", headers=headers)
    check("GET /orders/import/template -> 200", r.status_code == 200)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    template_cols = [str(c).strip() for c in next(ws.iter_rows(values_only=True))]
    expected_cols = [
        "order_group_id",
        "customer_id",
        "order_date",
        "delivery_method",
        "delivery_address",
        "payment_type",
        "product_id",
        "variant_id",
        "quantity",
        "unit_price",
    ]
    check(
        "Order template has EXACTLY 10 expected columns in order",
        template_cols == expected_cols,
        f"got: {template_cols}",
    )

    # 2. Multi-item orders grouped by order_group_id
    rows = [
        # Order 1 (2 items, takeaway, blank unit_price -> should use catalogue price)
        ["ORD-GRP-001", "9222222222", "2026-09-11", "takeaway", "", "credit", "CHR-STD-01", "", 4, ""],
        ["ORD-GRP-001", "9222222222", "2026-09-11", "takeaway", "", "credit", "DSK-PRO-02", "", 1, 7500.0],  # custom price
        # Order 2 (1 item, home delivery with address)
        ["ORD-GRP-002", cust1_id, "2026-09-11", "home_delivery", "99 MG Road, Suite 4", "cash", prod1_id, "", 2, 2400.0],
    ]
    xlsx_bytes = create_xlsx_bytes(expected_cols, rows)
    r = client.post(
        "/orders/import",
        headers=headers,
        files={"file": ("orders.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /orders/import grouped orders -> 200", r.status_code == 200, r.text)
    data = r.json()
    check("Created 2 orders (total 3 item rows)", data["success_count"] == 2 and data["total_rows"] == 3, str(data))

    # Verify DB orders and items
    db = SessionLocal()
    orders = db.query(SalesOrder).filter(SalesOrder.organization_id == org_id).order_by(SalesOrder.created_at.asc()).all()
    check("Found 2 sales orders in DB", len(orders) == 2)
    o1 = orders[0]
    check("Order 1 has 2 items", len(o1.items) == 2)
    # Check item 1 used catalogue fallback price (2500.0)
    item1 = next((i for i in o1.items if i.product_id == prod1_id), None)
    check("Order 1 item 1 used catalogue price 2500.0", item1 is not None and item1.unit_price == 2500.0)
    db.close()

    # 3. Invalid orders (conflicting customer in group, home delivery missing address on cust with no address)
    db = SessionLocal()
    cust_no_addr = Customer(organization_id=org_id, name="No Address Cust", is_active=True)
    db.add(cust_no_addr)
    db.commit()
    db.refresh(cust_no_addr)
    cust_no_addr_id = cust_no_addr.id
    db.close()

    bad_rows = [
        # Conflicting customer in same order_group_id
        ["ORD-BAD-01", cust1_id, "2026-09-11", "takeaway", "", "cash", "CHR-STD-01", "", 1, 2500.0],
        ["ORD-BAD-01", cust_no_addr_id, "2026-09-11", "takeaway", "", "cash", "CHR-STD-01", "", 1, 2500.0],
        # Home delivery with no address anywhere
        ["ORD-BAD-02", cust_no_addr_id, "2026-09-11", "home_delivery", "", "cash", "CHR-STD-01", "", 1, 2500.0],
    ]
    bad_xlsx = create_xlsx_bytes(expected_cols, bad_rows)
    r = client.post(
        "/orders/import",
        headers=headers,
        files={"file": ("bad_orders.xlsx", bad_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check("POST /orders/import bad rows -> 200 with structured errors", r.status_code == 200)
    res = r.json()
    check("Recorded errors for conflicting customer & missing address", res["error_count"] >= 2, str(res))


def main():
    print("==================================================")
    print("STARTING BULK IMPORT AUTOMATED VERIFICATION SUITE")
    print("==================================================")
    org_id, headers = register_org()

    run_customer_import_tests(org_id, headers)
    run_purchase_import_tests(org_id, headers)
    run_order_import_tests(org_id, headers)

    print("\n==================================================")
    print(f"RESULTS: {passed_count} PASSED, {failed_count} FAILED")
    print("==================================================")
    if failed_count > 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
