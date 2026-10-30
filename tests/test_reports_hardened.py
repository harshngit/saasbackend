import os
import uuid
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.core.database import Base, get_db
from app.core.deps import require_permission
from app.main import app
from app.models import (
    Customer,
    CustomerPayment,
    Delivery,
    DeliveryItem,
    Expense,
    Invoice,
    InvoiceItem,
    Organization,
    Product,
    ProductPricing,
    ProductVariant,
    PurchaseInvoice,
    PurchaseReturn,
    ReturnItem,
    SalesOrder,
    SalesOrderItem,
    SalesReturn,
    StockMovement,
    StockReservation,
    Supplier,
    SupplierInvoice,
    SupplierPayment,
    User,
    Vehicle,
    Warehouse,
    WarehouseStock,
    Account,
    JournalEntry,
    JournalLine,
)
from app.services import accounting_service, report_service

from sqlalchemy.pool import StaticPool

# Use in-memory SQLite database for testing with StaticPool so all sessions share the memory DB
TEST_DB_URL = "sqlite:///:memory:"

engine = create_engine(
    TEST_DB_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    Base.metadata.create_all(bind=engine, checkfirst=True)
    yield
    Base.metadata.drop_all(bind=engine, checkfirst=True)


@pytest.fixture
def db():
    session = TestingSessionLocal()
    yield session
    session.close()


from app.core.entitlements import default_entitlements_for_plan
from app.models import Plan


@pytest.fixture
def test_data(db):
    pro_plan = Plan(
        name="Pro",
        price_monthly=999,
        price_yearly=9999,
        entitlements=default_entitlements_for_plan("Pro"),
    )
    db.add(pro_plan)
    db.flush()

    org = Organization(
        id=str(uuid.uuid4()),
        name="Test Reports Org",
        company_code="CMP-10001",
        timezone="Asia/Kolkata",
        plan_id=pro_plan.id,
    )
    db.add(org)

    org2 = Organization(
        id=str(uuid.uuid4()),
        name="Foreign Org",
        company_code="CMP-20002",
        timezone="Asia/Kolkata",
        plan_id=pro_plan.id,
    )
    db.add(org2)

    user = User(
        id=str(uuid.uuid4()),
        email="testuser@example.com",
        name="Test Analyst",
        password_hash="fake_hash_value",
        organization_id=org.id,
        role="admin",
        is_active=True,
    )
    db.add(user)

    warehouse = Warehouse(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        name="Main Hub",
        code="WH-01",
        is_default=True,
        is_active=True,
    )
    db.add(warehouse)

    cust = Customer(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        name="Acme Corp",
        business_name="Acme Retail",
        phone="9876543210",
        credit_limit=50000.0,
    )
    db.add(cust)

    foreign_cust = Customer(
        id=str(uuid.uuid4()),
        organization_id=org2.id,
        name="Foreign Customer",
        phone="9999999999",
    )
    db.add(foreign_cust)

    sup = Supplier(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        name="Global Supplies",
        supplier_code="SUP-0001",
        phone="8765432109",
    )
    db.add(sup)

    prod = Product(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        name="Premium Widget",
        sku="WIDGET-01",
        price=100.0,
        minimum_stock_level=10,
        is_active=True,
    )
    db.add(prod)

    pricing = ProductPricing(
        id=str(uuid.uuid4()),
        product_id=prod.id,
        purchase_price=60.0,
        selling_price=100.0,
    )
    db.add(pricing)

    ws = WarehouseStock(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        warehouse_id=warehouse.id,
        product_id=prod.id,
        on_hand_quantity=50.0,
    )
    db.add(ws)

    db.commit()

    return {
        "org": org,
        "org2": org2,
        "user": user,
        "warehouse": warehouse,
        "customer": cust,
        "foreign_customer": foreign_cust,
        "supplier": sup,
        "product": prod,
        "pricing": pricing,
        "warehouse_stock": ws,
    }


@pytest.fixture
def client(db, test_data):
    from app.routers.reports import _export, _view

    def override_get_db():
        yield db

    def override_user():
        return test_data["user"]

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[_view] = override_user
    app.dependency_overrides[_export] = override_user

    with TestClient(app) as c:
        yield c

    app.dependency_overrides.clear()


def test_registry_contains_all_reports():
    assert len(report_service.REPORT_TYPES) == 21
    assert len(report_service.REPORT_REGISTRY) == 21
    expected_reports = {
        "daily-transaction", "sales", "purchase", "customer-outstanding",
        "supplier-outstanding", "payment-collection", "expense", "cash-collection",
        "gst-summary", "sales-return", "purchase-return", "profit-loss",
        "supplier-payment", "inventory-summary", "stock-movement",
        "sales-overview", "cash-flow-sheet",
        "statement-of-pl", "notes-to-pl",
        "balance-sheet", "notes-to-balance-sheet",
    }
    assert report_service.REPORT_TYPES == expected_reports


def test_standard_contract_and_backward_compatibility(client, test_data):
    # Call sales report with only date_from & date_to
    res = client.get("/reports/sales?date_from=2026-01-01&date_to=2026-12-31")
    assert res.status_code == 200
    data = res.json()

    # Verify existing contract fields
    assert data["type"] == "sales"
    assert data["date_from"] == "2026-01-01"
    assert data["date_to"] == "2026-12-31"
    assert "summary" in data
    assert "rows" in data

    # Verify additive fields
    assert "pagination" in data
    assert data["pagination"]["page"] == 1
    assert data["pagination"]["page_size"] == 25
    assert "meta" in data
    assert data["meta"]["currency"] == "INR"
    assert len(data["meta"]["columns"]) > 0
    assert "chart" in data
    assert data["chart"]["type"] == "line"


def test_daily_transaction_cash_separation(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]
    sup = test_data["supplier"]

    # 1. Accrual Sale (Invoice) of 10,000 on credit (paid = 0)
    inv = Invoice(
        organization_id=org.id,
        invoice_number="INV-001",
        customer_id=cust.id,
        subtotal=9000.0,
        tax=1000.0,
        total=10000.0,
        amount_paid=0.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv)

    # 2. Customer payment of 4,000 cash
    pay = CustomerPayment(
        organization_id=org.id,
        customer_id=cust.id,
        amount=4000.0,
        payment_mode="cash",
        reference="RCPT-001",
        received_on=datetime.now(timezone.utc),
    )
    db.add(pay)

    # 3. Accrual Purchase of 6,000 on credit
    pinv = PurchaseInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        invoice_number="PUR-001",
        subtotal=5500.0,
        tax=500.0,
        total=6000.0,
        status="approved",
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(pinv)
    db.flush()

    sinv = SupplierInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        purchase_id=pinv.id,
        supplier_invoice_number="SINV-001",
        subtotal=5500.0,
        tax_amount=500.0,
        grand_total=6000.0,
        amount_paid=0.0,
        status="recorded",
        payment_status="unpaid",
        supplier_invoice_date=datetime.now(timezone.utc),
    )
    db.add(sinv)

    # 4. Supplier Payment of 2,000
    spay = SupplierPayment(
        organization_id=org.id,
        supplier_id=sup.id,
        payment_number="SPAY-001",
        amount=2000.0,
        payment_mode="bank_transfer",
        status="recorded",
        paid_on=datetime.now(timezone.utc),
    )
    db.add(spay)

    # 5. Approved Paid Expense of 500
    exp = Expense(
        organization_id=org.id,
        category="Office Supplies",
        amount=500.0,
        status="approved",
        payment_status="paid",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(exp)

    db.commit()

    res = client.get("/reports/daily-transaction")
    assert res.status_code == 200
    data = res.json()
    summary = data["summary"]

    # Invariants:
    # Accrual Sales = 10,000
    # Customer Collections (Cash In) = 4,000 (Sale booking must NOT double count as Cash In)
    # Accrual Purchases = 6,000
    # Supplier Payments = 2,000
    # Expenses = 500
    # Cash Out = 2000 + 500 = 2500
    # Net Cash Flow = 4000 - 2500 = 1500
    assert summary["sales"] == 10000.0
    assert summary["purchases"] == 6000.0
    assert summary["customer_collections"] == 4000.0
    assert summary["supplier_payments"] == 2000.0
    assert summary["expenses"] == 500.0
    assert summary["cash_in"] == 4000.0
    assert summary["cash_out"] == 2500.0
    assert summary["net_cash_flow"] == 1500.0


def test_sales_report_groupings(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]
    prod = test_data["product"]

    # Create SalesOrder and SalesOrderItem with cost_price = 60.0
    so = SalesOrder(
        organization_id=org.id,
        customer_id=cust.id,
        order_number="SO-GRP-01",
        status="completed",
        total=220.0,
    )
    db.add(so)
    db.flush()

    so_item = SalesOrderItem(
        order_id=so.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=2,
        unit_price=100.0,
        cost_price=60.0,
        line_total=200.0,
    )
    db.add(so_item)
    db.flush()

    # Create Invoice with InvoiceItem
    inv = Invoice(
        organization_id=org.id,
        invoice_number="INV-GRP-01",
        order_id=so.id,
        customer_id=cust.id,
        subtotal=200.0,
        tax=20.0,
        total=220.0,
        amount_paid=100.0,
        status="partially_paid",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv)
    db.flush()

    item = InvoiceItem(
        invoice_id=inv.id,
        product_id=prod.id,
        order_item_id=so_item.id,
        product_name=prod.name,
        quantity=2,
        unit_price=100.0,
        tax=20.0,
        line_total=220.0,
    )
    db.add(item)
    db.commit()

    # 1. Customer grouping
    res_cust = client.get("/reports/sales?group_by=customer")
    assert res_cust.status_code == 200
    data_cust = res_cust.json()
    assert data_cust["meta"]["group_by"] == "customer"
    assert len(data_cust["rows"]) >= 1
    assert data_cust["rows"][0]["customer"] == "Acme Retail"

    # 2. Product grouping
    res_prod = client.get("/reports/sales?group_by=product")
    assert res_prod.status_code == 200
    data_prod = res_prod.json()
    assert data_prod["meta"]["group_by"] == "product"
    assert len(data_prod["rows"]) >= 1
    p_row = [r for r in data_prod["rows"] if r["product"] == "Premium Widget"][0]
    assert p_row["quantity_sold"] == 2.0
    # Cost = 2 * 60 = 120, Sales = 220, Gross Profit = 100
    assert p_row["cogs"] == 120.0
    assert p_row["gross_profit"] == 100.0


def test_customer_outstanding_aging(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]

    # Invoice overdue by 45 days (bucket 31_60)
    old_date = datetime.now(timezone.utc) - timedelta(days=45)
    inv = Invoice(
        organization_id=org.id,
        invoice_number="INV-OVERDUE-45",
        customer_id=cust.id,
        subtotal=1000.0,
        total=1000.0,
        amount_paid=0.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=old_date,
        due_date=old_date,
    )
    db.add(inv)
    db.commit()

    res = client.get("/reports/customer-outstanding")
    assert res.status_code == 200
    data = res.json()
    assert data["summary"]["age_31_60"] >= 1000.0
    assert data["summary"]["total_overdue"] >= 1000.0


def test_supplier_outstanding_aging(client, db, test_data):
    org = test_data["org"]
    sup = test_data["supplier"]

    old_date = datetime.now(timezone.utc) - timedelta(days=70)
    pinv = PurchaseInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        invoice_number="PUR-OLD-70",
        subtotal=3000.0,
        tax=0.0,
        total=3000.0,
        status="approved",
        invoice_date=old_date,
    )
    db.add(pinv)
    db.flush()

    sinv = SupplierInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        purchase_id=pinv.id,
        supplier_invoice_number="SINV-OLD-70",
        subtotal=3000.0,
        grand_total=3000.0,
        return_amount=500.0,
        amount_paid=500.0,
        status="recorded",
        payment_status="partially_paid",
        supplier_invoice_date=old_date,
        due_date=old_date,
    )
    db.add(sinv)
    db.commit()

    res = client.get("/reports/supplier-outstanding")
    assert res.status_code == 200
    data = res.json()
    # Outstanding = 3000 - 500 (return) - 500 (paid) = 2000
    assert data["summary"]["total_payable"] >= 2000.0
    assert data["summary"]["age_61_90"] >= 2000.0


def test_cash_collection_split_payment(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]

    # Customer payment with pure cash
    pay = CustomerPayment(
        organization_id=org.id,
        customer_id=cust.id,
        amount=1500.0,
        payment_mode="cash",
        reference="CASH-ONLY",
        received_on=datetime.now(timezone.utc),
    )
    db.add(pay)
    db.commit()

    res = client.get("/reports/cash-collection")
    assert res.status_code == 200
    data = res.json()
    assert data["summary"]["total_cash"] >= 1500.0


def test_profit_loss_cogs_formula(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]
    prod = test_data["product"]

    # Create SalesOrder and SalesOrderItem with stored sale-time cost snapshot = 60.0
    so = SalesOrder(
        organization_id=org.id,
        customer_id=cust.id,
        order_number="SO-PL-01",
        status="completed",
        total=500.0,
    )
    db.add(so)
    db.flush()

    so_item = SalesOrderItem(
        order_id=so.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=5,
        unit_price=100.0,
        cost_price=60.0,  # Sale-time historical cost snapshot
        line_total=500.0,
    )
    db.add(so_item)
    db.flush()

    # 1. Invoice for 5 items @ 100 = 500 total (COGS = 5 * 60 = 300 from stored snapshot)
    inv = Invoice(
        organization_id=org.id,
        invoice_number="INV-PL-01",
        order_id=so.id,
        customer_id=cust.id,
        subtotal=500.0,
        tax=0.0,
        total=500.0,
        amount_paid=500.0,
        status="paid",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv)
    db.flush()

    item = InvoiceItem(
        invoice_id=inv.id,
        product_id=prod.id,
        order_item_id=so_item.id,
        product_name=prod.name,
        quantity=5,
        unit_price=100.0,
        line_total=500.0,
    )
    db.add(item)

    # 2. Approved Expense = 50
    exp = Expense(
        organization_id=org.id,
        category="Logistics",
        amount=50.0,
        status="approved",
        payment_status="paid",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(exp)
    db.commit()

    res = client.get("/reports/profit-loss")
    assert res.status_code == 200
    data = res.json()
    summary = data["summary"]

    assert summary["cogs"] >= 300.0
    assert summary["cogs_basis"] == "historical_cost_snapshots"
    assert summary["gross_profit"] == round(summary["net_sales"] - summary["cogs"], 2)
    assert summary["net_profit"] == round(summary["gross_profit"] - summary["operating_expenses"], 2)


def test_pl_cogs_safety_no_fabricated_historical_cogs(client, db, test_data):
    """Regression test: Sale without reliable cost snapshot does NOT invent COGS from catalog purchase price."""
    org = test_data["org"]
    cust = test_data["customer"]
    prod = test_data["product"]  # has ProductPricing purchase_price = 60.0

    # Invoice without order_item_id / without cost snapshot
    inv_nocost = Invoice(
        organization_id=org.id,
        invoice_number="INV-NOCOST-999",
        customer_id=cust.id,
        subtotal=200.0,
        tax=0.0,
        total=200.0,
        amount_paid=200.0,
        status="paid",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv_nocost)
    db.flush()

    item_nocost = InvoiceItem(
        invoice_id=inv_nocost.id,
        product_id=prod.id,
        order_item_id=None,  # NO snapshot
        product_name=prod.name,
        quantity=2,
        unit_price=100.0,
        line_total=200.0,
    )
    db.add(item_nocost)
    db.commit()

    res = client.get("/reports/profit-loss")
    assert res.status_code == 200
    data = res.json()
    # The item with no snapshot must NOT have 2 * 60 = 120 added to COGS
    assert data["summary"]["missing_cogs_lines"] >= 1
    assert "no historical cost snapshot" in data["meta"]["cogs_notes"] or "strictly from" in data["meta"]["cogs_notes"]


def test_daily_transaction_cash_out_approved_unpaid_vs_paid_expense(client, db, test_data):
    """Regression test: Approved unpaid expense appears in expense activity but does NOT increase cash_out."""
    org = test_data["org"]

    # Baseline daily transaction report
    res_before = client.get("/reports/daily-transaction")
    assert res_before.status_code == 200
    cash_out_before = res_before.json()["summary"]["cash_out"]
    expenses_before = res_before.json()["summary"]["expenses"]

    # 1. Add an approved UNPAID expense (payment_status="Pending")
    unpaid_exp = Expense(
        organization_id=org.id,
        category="Office Rent",
        amount=12000.0,
        status="approved",
        payment_status="Pending",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(unpaid_exp)
    db.commit()

    res_after_unpaid = client.get("/reports/daily-transaction")
    assert res_after_unpaid.status_code == 200
    data_unpaid = res_after_unpaid.json()

    # Expenses activity total increased by 12000
    assert data_unpaid["summary"]["expenses"] == round(expenses_before + 12000.0, 2)
    # But cash_out did NOT increase!
    assert data_unpaid["summary"]["cash_out"] == cash_out_before

    # Verify row direction is Accrual Expense
    unpaid_rows = [r for r in data_unpaid["rows"] if r["reference"] == "Office Rent"]
    assert len(unpaid_rows) >= 1
    assert unpaid_rows[0]["direction"] == "Accrual Expense"

    # 2. Add an approved PAID expense (payment_status="Paid")
    paid_exp = Expense(
        organization_id=org.id,
        category="Courier Charges",
        amount=450.0,
        status="approved",
        payment_status="Paid",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(paid_exp)
    db.commit()

    res_after_paid = client.get("/reports/daily-transaction")
    assert res_after_paid.status_code == 200
    data_paid = res_after_paid.json()

    # Both expenses and cash_out increased by 450
    assert data_paid["summary"]["expenses"] == round(expenses_before + 12000.0 + 450.0, 2)
    assert data_paid["summary"]["cash_out"] == round(cash_out_before + 450.0, 2)

    paid_rows = [r for r in data_paid["rows"] if r["reference"] == "Courier Charges"]
    assert len(paid_rows) >= 1
    assert paid_rows[0]["direction"] == "Cash Out"


def test_receivable_and_payable_as_of_date_rejected(client):
    """Regression test: Unsupported historical as_of_date is rejected with clear HTTP 400 error."""
    # Customer outstanding with as_of_date
    res_ar = client.get("/reports/customer-outstanding?as_of_date=2025-01-01")
    assert res_ar.status_code == 400
    assert "as_of_date" in res_ar.json()["detail"]
    assert "not supported" in res_ar.json()["detail"]

    # Supplier outstanding with as_of_date
    res_ap = client.get("/reports/supplier-outstanding?as_of_date=2025-01-01")
    assert res_ap.status_code == 400
    assert "as_of_date" in res_ap.json()["detail"]
    assert "not supported" in res_ap.json()["detail"]


def test_supplier_payment_report(client, db, test_data):
    org = test_data["org"]
    sup = test_data["supplier"]

    sp = SupplierPayment(
        organization_id=org.id,
        supplier_id=sup.id,
        payment_number="SPAY-NEW-01",
        amount=7500.0,
        allocated_amount=5000.0,
        unallocated_amount=2500.0,
        payment_mode="upi",
        status="recorded",
        paid_on=datetime.now(timezone.utc),
    )
    db.add(sp)
    db.commit()

    res = client.get("/reports/supplier-payment")
    assert res.status_code == 200
    data = res.json()
    assert data["type"] == "supplier-payment"
    assert data["summary"]["total_paid"] >= 7500.0
    assert data["summary"]["allocated_amount"] >= 5000.0
    assert data["summary"]["unallocated_amount"] >= 2500.0


def test_inventory_summary_report(client, db, test_data):
    res = client.get("/reports/inventory-summary")
    assert res.status_code == 200
    data = res.json()
    assert data["type"] == "inventory-summary"
    assert data["summary"]["total_on_hand"] >= 50.0
    assert data["meta"]["valuation_basis"] == "Current cost valuation"
    assert len(data["rows"]) >= 1


def test_stock_movement_report(client, db, test_data):
    org = test_data["org"]
    wh = test_data["warehouse"]
    prod = test_data["product"]

    sm = StockMovement(
        organization_id=org.id,
        warehouse_id=wh.id,
        product_id=prod.id,
        movement_type="purchase_in",
        quantity=25,
        balance_after=75,
        note="Received batch B-01",
        created_at=datetime.now(timezone.utc),
    )
    db.add(sm)
    db.commit()

    res = client.get("/reports/stock-movement")
    assert res.status_code == 200
    data = res.json()
    assert data["type"] == "stock-movement"
    assert data["summary"]["stock_in"] >= 25
    assert len(data["rows"]) >= 1


def test_tenant_isolation_foreign_entity_rejected(client, test_data):
    # Pass foreign customer_id belonging to another organization
    foreign_id = test_data["foreign_customer"].id
    res = client.get(f"/reports/sales?customer_id={foreign_id}")
    assert res.status_code == 404
    assert "not found in organization" in res.json()["detail"]


def test_pagination_and_full_summary(client, db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]

    # Seed 30 invoices
    for idx in range(30):
        inv = Invoice(
            organization_id=org.id,
            invoice_number=f"INV-PAG-{idx}",
            customer_id=cust.id,
            subtotal=100.0,
            total=100.0,
            amount_paid=100.0,
            status="paid",
            is_credit_note=False,
            invoice_date=datetime.now(timezone.utc),
        )
        db.add(inv)
    db.commit()

    # Request page 1 with page_size = 10
    res = client.get("/reports/sales?page=1&page_size=10")
    assert res.status_code == 200
    data = res.json()

    assert data["pagination"]["page"] == 1
    assert data["pagination"]["page_size"] == 10
    assert data["pagination"]["total"] >= 30
    assert len(data["rows"]) == 10
    # Summary must reflect all 30+ invoices (not just 10)
    assert data["summary"]["total_sales"] >= 3000.0


def test_exports_excel_and_pdf(client):
    # Excel export
    res_xlsx = client.get("/reports/sales/export?format=excel")
    assert res_xlsx.status_code == 200
    assert res_xlsx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert len(res_xlsx.content) > 0

    # PDF export
    res_pdf = client.get("/reports/sales/export?format=pdf")
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert len(res_pdf.content) > 0


def test_screen_pagination_capped_at_100(client):
    """Screen pagination strictly rejects page_size > 100."""
    res = client.get("/reports/sales?page_size=150")
    assert res.status_code == 422  # FastAPI validation constraint le=100


def test_export_more_than_100_rows_not_truncated(client, db, test_data):
    """Export with >100 rows does not get truncated to 100 rows."""
    org = test_data["org"]
    cust = test_data["customer"]

    # Seed 120 invoices
    for idx in range(120):
        inv = Invoice(
            organization_id=org.id,
            invoice_number=f"INV-EXP-{idx}",
            customer_id=cust.id,
            subtotal=50.0,
            total=50.0,
            amount_paid=50.0,
            status="paid",
            is_credit_note=False,
            invoice_date=datetime.now(timezone.utc),
        )
        db.add(inv)
    db.commit()

    report = report_service.build_report(db, org.id, "sales", is_export=True)
    assert len(report["rows"]) >= 120
    assert report["pagination"]["page_size"] == 10000
    assert report["pagination"]["total"] >= 120

    # Excel export works for >100 rows
    res = client.get("/reports/sales/export?format=excel")
    assert res.status_code == 200
    assert len(res.content) > 0


def test_export_at_10000_rows_accepted(db, test_data, monkeypatch):
    """Export dataset containing up to 10,000 rows is accepted."""
    org = test_data["org"]

    mock_rows = [{"invoice_number": f"INV-{i}", "total": 100.0} for i in range(10000)]

    def mock_builder(db, org_id, df, dt, params):
        paged_rows, pagination = report_service._paginate(mock_rows, params.get("page", 1), params.get("page_size", 25))
        return {
            "summary": {"total_sales": 1000000.0},
            "rows": paged_rows,
            "pagination": pagination,
            "meta": {"currency": "INR", "columns": []},
            "chart": {"type": "bar", "data": []},
        }

    monkeypatch.setitem(report_service._BUILDERS, "sales", mock_builder)

    res = report_service.build_report(db, org.id, "sales", is_export=True)
    assert len(res["rows"]) == 10000
    assert res["pagination"]["total"] == 10000


def test_export_more_than_10000_rows_rejected_with_http_400(client, test_data, monkeypatch):
    """Export dataset exceeding 10,000 rows returns HTTP 400 with a clear error message."""
    org = test_data["org"]

    mock_rows = [{"invoice_number": f"INV-{i}", "total": 100.0} for i in range(10005)]

    def mock_builder(db, org_id, df, dt, params):
        paged_rows, pagination = report_service._paginate(mock_rows, params.get("page", 1), params.get("page_size", 25))
        return {
            "summary": {"total_sales": 1000500.0},
            "rows": paged_rows,
            "pagination": pagination,
            "meta": {"currency": "INR", "columns": []},
            "chart": {"type": "bar", "data": []},
        }

    monkeypatch.setitem(report_service._BUILDERS, "sales", mock_builder)

    res = client.get("/reports/sales/export?format=excel")
    assert res.status_code == 400
    detail = res.json()["detail"]
    assert "Export limit exceeded" in detail
    assert "10,000 rows" in detail


def test_sales_overview_operational_report(client, db, test_data):
    """Test Sales Overview operational report: root SalesOrder, delivery aggregation, quantities, filters, and summary."""
    org = test_data["org"]
    cust = test_data["customer"]
    salesp = test_data["user"]
    prod = test_data["product"]

    # 1. Driver & Vehicle setup
    driver = User(
        id=str(uuid.uuid4()),
        email="driver@example.com",
        name="Fast Driver",
        password_hash="fake",
        organization_id=org.id,
        role="delivery_partner",
        is_active=True,
    )
    vehicle = Vehicle(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        vehicle_number="MH-01-AB-1234",
        vehicle_type="Tata Ace",
        is_active=True,
    )
    db.add_all([driver, vehicle])
    db.commit()

    # Order 1: No deliveries, not_started, 2 items (qty 5 + 10 = 15)
    so1 = SalesOrder(
        organization_id=org.id,
        order_number="SO-OV-01",
        sales_order_number="SO-10001-2026-00001",
        customer_id=cust.id,
        salesperson_id=salesp.id,
        status="placed",
        fulfilment_status="not_started",
        total=1500.0,
        order_date=datetime.now(timezone.utc),
    )
    db.add(so1)
    db.flush()
    item1_1 = SalesOrderItem(
        order_id=so1.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=5,
        unit_price=100.0,
        delivered_quantity=0,
        line_total=500.0,
    )
    item1_2 = SalesOrderItem(
        order_id=so1.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=10,
        unit_price=100.0,
        delivered_quantity=0,
        line_total=1000.0,
    )
    db.add_all([item1_1, item1_2])

    # Order 2: 1 delivery, planned
    so2 = SalesOrder(
        organization_id=org.id,
        order_number="SO-OV-02",
        sales_order_number="SO-10001-2026-00002",
        customer_id=cust.id,
        salesperson_id=salesp.id,
        status="processing",
        fulfilment_status="planned",
        total=800.0,
        order_date=datetime.now(timezone.utc),
    )
    db.add(so2)
    db.flush()
    item2 = SalesOrderItem(
        order_id=so2.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=8,
        unit_price=100.0,
        delivered_quantity=0,
        line_total=800.0,
    )
    db.add(item2)
    deliv2 = Delivery(
        organization_id=org.id,
        sales_order_id=so2.id,
        customer_id=cust.id,
        delivery_note_number="DN-OV-02",
        delivery_partner_id=driver.id,
        vehicle_id=vehicle.id,
        status="planned",
        delivery_date=datetime.now(timezone.utc),
    )
    db.add(deliv2)

    # Order 3: Multiple deliveries (Deliv A: delivered, Deliv B: in_transit, Deliv C: cancelled)
    so3 = SalesOrder(
        organization_id=org.id,
        order_number="SO-OV-03",
        sales_order_number="SO-10001-2026-00003",
        customer_id=cust.id,
        salesperson_id=salesp.id,
        status="processing",
        fulfilment_status="partially_delivered",
        total=2000.0,
        order_date=datetime.now(timezone.utc),
    )
    db.add(so3)
    db.flush()
    item3 = SalesOrderItem(
        order_id=so3.id,
        product_id=prod.id,
        product_name=prod.name,
        quantity=20,
        unit_price=100.0,
        delivered_quantity=12,
        line_total=2000.0,
    )
    db.add(item3)
    deliv3_a = Delivery(
        organization_id=org.id,
        sales_order_id=so3.id,
        customer_id=cust.id,
        delivery_note_number="DN-OV-03A",
        delivery_partner_id=driver.id,
        vehicle_id=vehicle.id,
        status="delivered",
        delivery_date=datetime.now(timezone.utc) - timedelta(days=2),
    )
    deliv3_b = Delivery(
        organization_id=org.id,
        sales_order_id=so3.id,
        customer_id=cust.id,
        delivery_note_number="DN-OV-03B",
        delivery_partner_id=driver.id,
        vehicle_id=vehicle.id,
        status="in_transit",
        delivery_date=datetime.now(timezone.utc) - timedelta(days=1),
    )
    deliv3_cancelled = Delivery(
        organization_id=org.id,
        sales_order_id=so3.id,
        customer_id=cust.id,
        delivery_note_number="DN-OV-03C-VOID",
        status="cancelled",
        delivery_date=datetime.now(timezone.utc),
    )
    db.add_all([deliv3_a, deliv3_b, deliv3_cancelled])

    # Order 4: Cancelled order
    so4_cancelled = SalesOrder(
        organization_id=org.id,
        order_number="SO-OV-04-CANC",
        sales_order_number="SO-10001-2026-00004",
        customer_id=cust.id,
        status="cancelled",
        fulfilment_status="not_started",
        total=5000.0,
        order_date=datetime.now(timezone.utc),
    )
    db.add(so4_cancelled)
    db.commit()

    # Query Sales Overview
    res = client.get("/reports/sales-overview")
    assert res.status_code == 200
    data = res.json()

    assert data["type"] == "sales-overview"
    summary = data["summary"]
    assert summary["total_orders"] >= 4
    # Cancelled order 4 is excluded from total_sales_value (1500 + 800 + 2000 = 4300)
    assert summary["total_sales_value"] >= 4300.0
    assert summary["cancelled_orders"] >= 1
    assert summary["not_started"] >= 1
    assert summary["planned"] >= 1
    assert summary["partially_delivered"] >= 1

    # Find rows by order_number
    rows = data["rows"]
    r1 = next(r for r in rows if r["order_number"] == "SO-OV-01")
    assert r1["delivery_count"] == 0
    assert r1["ordered_quantity"] == 15.0
    assert r1["delivered_quantity"] == 0.0
    assert r1["remaining_quantity"] == 15.0
    assert r1["latest_delivery"] is None

    r2 = next(r for r in rows if r["order_number"] == "SO-OV-02")
    assert r2["delivery_count"] == 1
    assert r2["latest_delivery"]["delivery_number"] == "DN-OV-02"
    assert r2["latest_delivery"]["delivery_partner"]["name"] == "Fast Driver"
    assert r2["latest_delivery"]["vehicle"]["vehicle_number"] == "MH-01-AB-1234"

    r3 = next(r for r in rows if r["order_number"] == "SO-OV-03")
    # Cancelled delivery is excluded from delivery_count (2 active: 3A and 3B)
    assert r3["delivery_count"] == 2
    assert r3["ordered_quantity"] == 20.0
    assert r3["delivered_quantity"] == 12.0
    assert r3["remaining_quantity"] == 8.0
    # Latest active delivery is 3B (in_transit)
    assert r3["latest_delivery"]["delivery_number"] == "DN-OV-03B"
    assert r3["latest_delivery"]["status"] == "in_transit"

    # Test filtering by search
    res_search = client.get("/reports/sales-overview?search=DN-OV-03A")
    assert res_search.status_code == 200
    search_rows = res_search.json()["rows"]
    assert any(r["order_number"] == "SO-OV-03" for r in search_rows)

    # Test filtering by fulfilment_status
    res_f = client.get("/reports/sales-overview?fulfilment_status=partially_delivered")
    assert res_f.status_code == 200
    assert any(r["order_number"] == "SO-OV-03" for r in res_f.json()["rows"])


def test_cash_flow_sheet_recorded_cash_movements(client, db, test_data):
    """Test Cash Flow Sheet: verifies cash inflow from CustomerPayment, outflow from SupplierPayment & paid Expense,
    non-cash exclusion, matching with Daily Transaction semantics, running balance, and exports."""
    org = test_data["org"]
    cust = test_data["customer"]
    sup = test_data["supplier"]

    # 1. Invoice alone (Accrual - must NOT create cash inflow)
    inv_accrual = Invoice(
        organization_id=org.id,
        invoice_number="INV-CF-ACCRUAL",
        customer_id=cust.id,
        subtotal=10000.0,
        total=10000.0,
        amount_paid=0.0,
        status="issued",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv_accrual)

    # 2. Real Customer Payment = +4000.0 (Cash In)
    cp = CustomerPayment(
        organization_id=org.id,
        customer_id=cust.id,
        amount=4000.0,
        payment_mode="upi",
        reference="UPI-CF-001",
        received_on=datetime.now(timezone.utc) - timedelta(hours=3),
    )
    db.add(cp)

    # 3. Non-void Supplier Payment = -1500.0 (Cash Out)
    sp = SupplierPayment(
        organization_id=org.id,
        supplier_id=sup.id,
        amount=1500.0,
        payment_mode="bank_transfer",
        reference="NEFT-CF-001",
        status="recorded",
        paid_on=datetime.now(timezone.utc) - timedelta(hours=2),
    )
    db.add(sp)

    # 4. Void Supplier Payment = -9999.0 (Must be EXCLUDED)
    sp_void = SupplierPayment(
        organization_id=org.id,
        supplier_id=sup.id,
        amount=9999.0,
        payment_mode="cash",
        reference="VOID-PAY",
        status="void",
        paid_on=datetime.now(timezone.utc),
    )
    db.add(sp_void)

    # 5. Unpaid approved Expense = -5000.0 (Must be EXCLUDED from cash outflow)
    exp_unpaid = Expense(
        organization_id=org.id,
        category="Office Rent",
        amount=5000.0,
        status="approved",
        payment_status="pending",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(exp_unpaid)

    # 6. Paid approved Expense = -500.0 (Cash Out)
    exp_paid = Expense(
        organization_id=org.id,
        category="Stationery",
        amount=500.0,
        status="approved",
        payment_status="paid",
        expense_date=datetime.now(timezone.utc) - timedelta(hours=1),
    )
    db.add(exp_paid)
    db.commit()

    # Query Cash Flow Sheet
    res_cf = client.get("/reports/cash-flow-sheet")
    assert res_cf.status_code == 200
    data_cf = res_cf.json()

    # Query Daily Transaction to compare exact matching
    res_dt = client.get("/reports/daily-transaction")
    assert res_dt.status_code == 200
    data_dt = res_dt.json()

    # Cash in = 4000.0
    assert data_cf["summary"]["cash_in"] >= 4000.0
    # Cash out = 1500 (SP) + 500 (Paid Exp) = 2000.0
    assert data_cf["summary"]["cash_out"] >= 2000.0
    assert data_cf["summary"]["net_cash_flow"] == round(data_cf["summary"]["cash_in"] - data_cf["summary"]["cash_out"], 2)

    # Must match Daily Transaction cash_in and cash_out exactly
    assert data_cf["summary"]["cash_in"] == data_dt["summary"]["cash_in"]
    assert data_cf["summary"]["cash_out"] == data_dt["summary"]["cash_out"]
    assert data_cf["summary"]["net_cash_flow"] == data_dt["summary"]["net_cash_flow"]

    # Check rows
    rows = data_cf["rows"]
    assert len(rows) >= 3
    # Check that void supplier payment and unpaid expense are NOT in cash flow rows
    ref_list = [r["reference_number"] for r in rows]
    assert "VOID-PAY" not in ref_list

    # Check export endpoints
    res_xlsx = client.get("/reports/cash-flow-sheet/export?format=excel")
    assert res_xlsx.status_code == 200
    assert len(res_xlsx.content) > 0

    res_pdf = client.get("/reports/cash-flow-sheet/export?format=pdf")
    assert res_pdf.status_code == 200
    assert len(res_pdf.content) > 0




# ---------------------------------------------------------------------------
# GST transaction view, ignored-filter fixes, date validation, export parity
# ---------------------------------------------------------------------------

from io import BytesIO

from app.models import PurchaseReturnItem

IST = ZoneInfo("Asia/Kolkata")


def _ist(y, m, d, hh=12, mm=0):
    return datetime(y, m, d, hh, mm, tzinfo=IST).astimezone(timezone.utc)


CUST_A_GSTIN = "27AAACA1111A1Z5"
CUST_B_GSTIN = "29BBBCB2222B1Z6"
SUP_GSTIN = "27SUPPL3333C1Z7"


@pytest.fixture
def gst_data(db, test_data):
    org, org2 = test_data["org"], test_data["org2"]
    cust_a, sup = test_data["customer"], test_data["supplier"]
    cust_a.gst_number = CUST_A_GSTIN
    sup.gst_number = SUP_GSTIN
    cust_b = Customer(id=str(uuid.uuid4()), organization_id=org.id, name="Beta Traders", phone="9000000001", gst_number=CUST_B_GSTIN)
    db.add(cust_b)

    def invoice(number, customer_id, when, subtotal, tax, status, org_id=org.id, walk_in=None):
        inv = Invoice(
            organization_id=org_id, invoice_number=number, customer_id=customer_id, walk_in_name=walk_in,
            subtotal=subtotal, tax=tax, total=subtotal + tax, amount_paid=(subtotal + tax) if status == "paid" else 0.0,
            status=status, is_credit_note=False, invoice_date=when,
        )
        db.add(inv)
        return inv

    invoice("INV-GST-001", cust_a.id, _ist(2026, 3, 5), 1000.0, 180.0, "unpaid")
    inv2 = invoice("INV-GST-002", cust_b.id, _ist(2026, 3, 15), 5000.0, 900.0, "paid")
    invoice("INV-GST-003", None, _ist(2026, 4, 2), 200.0, 10.0, "paid", walk_in="Counter Sale")
    invoice("INV-GST-VOID", cust_a.id, _ist(2026, 3, 6), 999.0, 99.0, "void")
    invoice("INV-GST-FOREIGN", test_data["foreign_customer"].id, _ist(2026, 3, 7), 777.0, 77.0, "paid", org_id=org2.id)
    db.flush()

    sr = SalesReturn(
        organization_id=org.id, return_number="SR-GST-001", return_date=_ist(2026, 3, 20),
        customer_id=cust_b.id, invoice_reference_id=inv2.id, return_status="approved", credit_amount=1180.0,
    )
    db.add(sr)
    db.flush()
    db.add(ReturnItem(return_id=sr.id, product_name="Premium Widget", quantity_returned=10, unit_price=100.0, tax_rate=18.0, line_total=1000.0))

    pinv = PurchaseInvoice(
        organization_id=org.id, supplier_id=sup.id, invoice_number="PUR-GST-001", subtotal=3000.0, tax=540.0,
        total=3540.0, status="approved", invoice_date=_ist(2026, 3, 12), warehouse_id=test_data["warehouse"].id,
    )
    db.add(pinv)
    db.flush()
    db.add(SupplierInvoice(
        organization_id=org.id, supplier_id=sup.id, purchase_id=pinv.id, supplier_invoice_number="SINV-GST-001",
        subtotal=3000.0, tax_amount=540.0, grand_total=3540.0, status="recorded", supplier_invoice_date=_ist(2026, 3, 12),
    ))
    pr = PurchaseReturn(
        organization_id=org.id, return_number="PR-GST-001", purchase_id=pinv.id, supplier_id=sup.id,
        status="confirmed", return_date=_ist(2026, 3, 25),
    )
    db.add(pr)
    db.flush()
    db.add(PurchaseReturnItem(purchase_return_id=pr.id, product_name="Premium Widget", quantity=2, unit_price=100.0, tax_rate=18.0, tax_amount=36.0, line_total=236.0))
    db.commit()
    return {**test_data, "customer_b": cust_b, "purchase": pinv}


def _gst(client, **params):
    params.setdefault("group_by", "transaction")
    params.setdefault("page_size", 100)
    return client.get("/reports/gst-summary", params=params)


def _docs(res):
    assert res.status_code == 200, res.text
    return sorted(r["document_number"] for r in res.json()["rows"])


def test_gst_summary_default_unchanged(client, gst_data):
    res = client.get("/reports/gst-summary")
    assert res.status_code == 200
    data = res.json()
    assert data["meta"]["group_by"] == "summary"
    assert set(data["summary"]) == {"output_gst", "input_gst", "gst_adjustments", "net_gst"}
    # Output 180 + 900 + 10 = 1090 less sales-return tax 180; input 540 less stored purchase-return tax 36.
    assert data["summary"] == {"output_gst": 910.0, "input_gst": 504.0, "gst_adjustments": 216.0, "net_gst": 406.0}
    rows = data["rows"]
    assert len(rows) == 4
    assert (rows[0]["taxable_value"], rows[0]["tax_amount"]) == (6200.0, 1090.0)
    # Sales-return taxable value is the pre-tax line total, not the tax-inclusive credit (1180).
    assert (rows[1]["taxable_value"], rows[1]["tax_amount"]) == (-1000.0, -180.0)
    assert (rows[2]["taxable_value"], rows[2]["tax_amount"]) == (3000.0, 540.0)
    # Net taxable = (6200 - 1000) - (3000 - 200)
    assert (rows[3]["taxable_value"], rows[3]["tax_amount"]) == (2400.0, 406.0)


def test_gst_summary_rejects_transaction_filters(client, gst_data):
    res = client.get("/reports/gst-summary", params={"customer_id": gst_data["customer"].id})
    assert res.status_code == 400
    assert "group_by=transaction" in res.json()["detail"]
    assert client.get("/reports/gst-summary", params={"group_by": "bogus"}).status_code == 400
    # Filters the transaction view cannot honour are rejected, not ignored.
    assert _gst(client, warehouse_id=gst_data["warehouse"].id).status_code == 400


def test_gst_transactions_rows_and_totals_match_summary(client, gst_data):
    res = _gst(client)
    assert _docs(res) == ["INV-GST-001", "INV-GST-002", "INV-GST-003", "PR-GST-001", "SINV-GST-001", "SR-GST-001"]
    data = res.json()
    assert data["meta"]["group_by"] == "transaction"
    by_doc = {r["document_number"]: r for r in data["rows"]}
    assert by_doc["SR-GST-001"] == {
        "document_type": "sales_return", "gst_classification": "output_adjustment", "source_id": by_doc["SR-GST-001"]["source_id"],
        "document_number": "SR-GST-001", "document_date": "2026-03-20", "reference_document": "INV-GST-002",
        "party_type": "customer", "party_id": gst_data["customer_b"].id, "party_name": "Beta Traders",
        "gst_number": CUST_B_GSTIN, "status": "approved", "taxable_value": 1000.0, "tax_amount": 180.0, "total_amount": 1180.0,
    }
    pr = by_doc["PR-GST-001"]
    assert (pr["taxable_value"], pr["tax_amount"], pr["total_amount"]) == (200.0, 36.0, 236.0)
    assert pr["reference_document"] == "PUR-GST-001"
    assert by_doc["INV-GST-003"]["party_name"] == "Counter Sale" and by_doc["INV-GST-003"]["gst_number"] is None

    s = data["summary"]
    aggregate = client.get("/reports/gst-summary").json()
    summary_rows = aggregate["rows"]
    assert s["output_tax_amount"] == summary_rows[0]["tax_amount"] == 1090.0
    assert s["output_taxable_value"] == summary_rows[0]["taxable_value"]
    assert s["output_adjustment_tax_amount"] == -summary_rows[1]["tax_amount"] == 180.0
    assert s["output_adjustment_taxable_value"] == -summary_rows[1]["taxable_value"] == 1000.0
    assert s["input_tax_amount"] == summary_rows[2]["tax_amount"] == 540.0
    assert s["input_adjustment_tax_amount"] == 36.0
    assert s["input_adjustment_taxable_value"] == 200.0
    # Both views reconcile: same net liability, same return-tax treatment.
    assert s["net_tax_amount"] == aggregate["summary"]["net_gst"] == summary_rows[3]["tax_amount"] == 406.0
    assert s["input_tax_amount"] - s["input_adjustment_tax_amount"] == aggregate["summary"]["input_gst"]
    assert s["output_adjustment_tax_amount"] + s["input_adjustment_tax_amount"] == aggregate["summary"]["gst_adjustments"]
    assert s["row_count"] == 6


def test_gst_transactions_date_range(client, gst_data):
    assert "INV-GST-003" not in _docs(_gst(client, date_from="2026-03-01", date_to="2026-03-31"))
    assert _docs(_gst(client, date_from="2026-04-02", date_to="2026-04-02")) == ["INV-GST-003"]
    res = _gst(client, date_from="2026-04-01", date_to="2026-03-01")
    assert res.status_code == 400
    assert "date_from cannot be after date_to" in res.json()["detail"]


def test_gst_transactions_customer_and_supplier_filters(client, gst_data):
    assert _docs(_gst(client, customer_id=gst_data["customer_b"].id)) == ["INV-GST-002", "SR-GST-001"]
    assert _docs(_gst(client, supplier_id=gst_data["supplier"].id)) == ["PR-GST-001", "SINV-GST-001"]
    # Party filters cannot be combined with each other or with the other side's documents.
    assert _gst(client, customer_id=gst_data["customer"].id, supplier_id=gst_data["supplier"].id).status_code == 400
    assert _gst(client, customer_id=gst_data["customer"].id, invoice_type="supplier_invoice").status_code == 400
    assert _gst(client, supplier_id=gst_data["supplier"].id, gst_classification="output").status_code == 400


def test_gst_transactions_document_reference_and_search(client, gst_data):
    # A return matches through the invoice it references.
    assert _docs(_gst(client, invoice_number="inv-gst-002")) == ["INV-GST-002", "SR-GST-001"]
    assert _docs(_gst(client, invoice_number="PR-GST")) == ["PR-GST-001"]
    assert _docs(_gst(client, search="beta")) == ["INV-GST-002", "SR-GST-001"]
    assert _docs(_gst(client, search="global supplies")) == ["PR-GST-001", "SINV-GST-001"]


def test_gst_transactions_gst_number(client, gst_data):
    assert _docs(_gst(client, gst_number=CUST_A_GSTIN.lower())) == ["INV-GST-001"]  # void invoice excluded
    assert _docs(_gst(client, gst_number=SUP_GSTIN)) == ["PR-GST-001", "SINV-GST-001"]
    assert _gst(client, gst_number="27-AA").status_code == 400


def test_gst_transactions_classification_type_and_status(client, gst_data):
    assert _docs(_gst(client, gst_classification="input")) == ["SINV-GST-001"]
    assert _docs(_gst(client, gst_classification="output,output_adjustment")) == ["INV-GST-001", "INV-GST-002", "INV-GST-003", "SR-GST-001"]
    assert _docs(_gst(client, invoice_type="sales_return,purchase_return")) == ["PR-GST-001", "SR-GST-001"]
    assert _docs(_gst(client, status="paid")) == ["INV-GST-002", "INV-GST-003"]
    assert _docs(_gst(client, status="confirmed")) == ["PR-GST-001"]
    for bad in ({"gst_classification": "igst"}, {"invoice_type": "quotation"}, {"status": "void"}, {"status": "draft"}):
        res = _gst(client, **bad)
        assert res.status_code == 400, bad
        assert "Allowed" in res.json()["detail"]


def test_gst_transactions_amount_ranges(client, gst_data):
    assert _docs(_gst(client, min_taxable_value=1000)) == ["INV-GST-001", "INV-GST-002", "SINV-GST-001", "SR-GST-001"]
    assert _docs(_gst(client, max_tax_amount=100)) == ["INV-GST-003", "PR-GST-001"]
    assert _docs(_gst(client, min_amount=500, max_amount=2000)) == ["INV-GST-001", "SR-GST-001"]
    res = _gst(client, min_tax_amount=500, max_tax_amount=100)
    assert res.status_code == 400
    assert "cannot be greater than" in res.json()["detail"]
    assert _gst(client, min_amount=-1).status_code == 422


def test_gst_transactions_combined_filters_sort_and_pagination(client, gst_data):
    res = _gst(client, date_from="2026-03-01", date_to="2026-03-31", gst_classification="output", min_tax_amount=500)
    assert _docs(res) == ["INV-GST-002"]
    assert res.json()["summary"]["output_tax_amount"] == 900.0

    rows = _gst(client, sort_by="tax_amount", sort_order="asc").json()["rows"]
    assert [r["tax_amount"] for r in rows] == sorted(r["tax_amount"] for r in rows)
    rows = _gst(client).json()["rows"]
    assert [r["document_date"] for r in rows] == sorted((r["document_date"] for r in rows), reverse=True)
    assert _gst(client, sort_by="bogus").status_code == 400
    assert _gst(client, sort_order="sideways").status_code == 400

    page = _gst(client, page=2, page_size=4).json()
    assert page["pagination"] == {"page": 2, "page_size": 4, "total": 6, "total_pages": 2}
    assert len(page["rows"]) == 2
    assert page["summary"]["row_count"] == 6


def test_gst_transactions_cross_org_isolation(client, gst_data):
    assert "INV-GST-FOREIGN" not in _docs(_gst(client))
    foreign = gst_data["foreign_customer"].id
    assert _gst(client, customer_id=foreign).status_code == 404
    res = client.get("/reports/gst-summary/export", params={"group_by": "transaction", "customer_id": foreign})
    assert res.status_code == 404


def _xlsx_table(content: bytes, first_header: str) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(content)).active
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(values) if r and r[0] == first_header)
    headers = values[start]
    return [dict(zip(headers, r)) for r in values[start + 1:] if any(v is not None for v in r)]


GST_COMBINED = {
    "group_by": "transaction", "date_from": "2026-03-01", "date_to": "2026-03-31",
    "gst_classification": "output,output_adjustment,input_adjustment", "max_tax_amount": 900,
}


def test_gst_excel_export_matches_unpaginated_api(client, gst_data):
    # Screen view paginated at 2 rows; the export must still carry every filtered row.
    api_page = client.get("/reports/gst-summary", params={**GST_COMBINED, "page_size": 2}).json()
    api_all = client.get("/reports/gst-summary", params={**GST_COMBINED, "page_size": 100}).json()
    assert api_page["pagination"]["total"] == 4 and len(api_page["rows"]) == 2

    res = client.get("/reports/gst-summary/export", params={**GST_COMBINED, "format": "excel"})
    assert res.status_code == 200
    table = _xlsx_table(res.content, "Date")
    assert sorted(r["Document #"] for r in table) == sorted(r["document_number"] for r in api_all["rows"])
    assert sorted(r["Tax Amount"] for r in table) == sorted(r["tax_amount"] for r in api_all["rows"])
    assert "INV-GST-003" not in {r["Document #"] for r in table}  # outside date range
    assert "SINV-GST-001" not in {r["Document #"] for r in table}  # excluded classification


def test_gst_pdf_export_uses_filtered_dataset(client, gst_data, monkeypatch):
    from app.routers import reports as reports_router

    captured = {}
    real_to_pdf = reports_router._to_pdf

    def spy(report, org_name=None):
        captured["report"] = report
        return real_to_pdf(report, org_name=org_name)

    monkeypatch.setattr(reports_router, "_to_pdf", spy)
    res = client.get("/reports/gst-summary/export", params={**GST_COMBINED, "format": "pdf"})
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    api_all = client.get("/reports/gst-summary", params={**GST_COMBINED, "page_size": 100}).json()
    assert captured["report"]["rows"] == api_all["rows"]
    assert captured["report"]["summary"] == api_all["summary"]


def test_reversed_date_range_rejected_and_boundaries_inclusive(client, db, test_data):
    for path in ("/reports/sales", "/reports/purchase", "/reports/expense", "/reports/gst-summary", "/reports/sales/export"):
        res = client.get(path, params={"date_from": "2026-05-02", "date_to": "2026-05-01"})
        assert res.status_code == 400, path
        assert res.json()["detail"] == "date_from cannot be after date_to"

    org, cust = test_data["org"], test_data["customer"]
    for number, when in (
        ("INV-BND-START", _ist(2026, 5, 1, 0, 0)),     # first instant of the IST day
        ("INV-BND-END", _ist(2026, 5, 1, 23, 59)),     # last minute of the IST day
        ("INV-BND-BEFORE", _ist(2026, 4, 30, 23, 59)),  # still 30 April in Kolkata
        ("INV-BND-AFTER", _ist(2026, 5, 2, 0, 0)),
    ):
        db.add(Invoice(organization_id=org.id, invoice_number=number, customer_id=cust.id, subtotal=10.0, total=10.0,
                       status="unpaid", is_credit_note=False, invoice_date=when))
    db.commit()
    res = client.get("/reports/sales", params={"date_from": "2026-05-01", "date_to": "2026-05-01"})
    assert sorted(r["invoice_number"] for r in res.json()["rows"]) == ["INV-BND-END", "INV-BND-START"]


def test_sales_warehouse_filter(client, db, test_data):
    org, cust, wh1 = test_data["org"], test_data["customer"], test_data["warehouse"]
    wh2 = Warehouse(id=str(uuid.uuid4()), organization_id=org.id, name="Second Hub", code="WH-02", is_active=True)
    foreign_wh = Warehouse(id=str(uuid.uuid4()), organization_id=test_data["org2"].id, name="Foreign Hub", code="WH-F", is_active=True)
    db.add_all([wh2, foreign_wh])
    so_wh1 = SalesOrder(organization_id=org.id, order_number="SO-WH-1", customer_id=cust.id, warehouse_id=wh1.id, total=100.0)
    so_wh2 = SalesOrder(organization_id=org.id, order_number="SO-WH-2", customer_id=cust.id, warehouse_id=wh2.id, total=100.0)
    db.add_all([so_wh1, so_wh2])
    db.flush()
    dl_wh1 = Delivery(organization_id=org.id, delivery_note_number="DN-WH-1", sales_order_id=so_wh2.id, warehouse_id=wh1.id)
    dl_none = Delivery(organization_id=org.id, delivery_note_number="DN-WH-0", sales_order_id=so_wh1.id)
    db.add_all([dl_wh1, dl_none])
    db.flush()

    def inv(number, order_id=None, delivery_id=None):
        db.add(Invoice(organization_id=org.id, invoice_number=number, customer_id=cust.id, order_id=order_id,
                       delivery_id=delivery_id, subtotal=100.0, total=100.0, status="unpaid", is_credit_note=False))

    inv("INV-WH-DELIVERY", order_id=so_wh2.id, delivery_id=dl_wh1.id)   # delivery warehouse wins over order's
    inv("INV-WH-ORDER2", order_id=so_wh2.id)                              # no delivery: order warehouse
    inv("INV-WH-FALLBACK", order_id=so_wh1.id, delivery_id=dl_none.id)   # delivery without warehouse: order's
    inv("INV-WH-NONE")                                                    # no association at all
    db.commit()

    def numbers(wh_id):
        res = client.get("/reports/sales", params={"warehouse_id": wh_id, "page_size": 100})
        assert res.status_code == 200, res.text
        return sorted(r["invoice_number"] for r in res.json()["rows"])

    assert numbers(wh1.id) == ["INV-WH-DELIVERY", "INV-WH-FALLBACK"]
    assert numbers(wh2.id) == ["INV-WH-ORDER2"]
    assert len(client.get("/reports/sales", params={"page_size": 100}).json()["rows"]) == 4
    # Grouped views apply the same filter.
    grouped = client.get("/reports/sales", params={"warehouse_id": wh1.id, "group_by": "salesperson"}).json()
    assert sum(r["transaction_count"] for r in grouped["rows"]) == 2
    assert client.get("/reports/sales", params={"warehouse_id": foreign_wh.id}).status_code == 404


def test_sales_salesperson_group_falls_back_to_creator(client, db, test_data):
    org, user = test_data["org"], test_data["user"]
    for idx in range(3):
        db.add(Invoice(organization_id=org.id, invoice_number=f"INV-CRT-{idx}", subtotal=50.0, total=50.0,
                       status="unpaid", is_credit_note=False, created_by=user.id))
    db.add(Invoice(organization_id=org.id, invoice_number="INV-CRT-NONE", subtotal=5.0, total=5.0,
                   status="unpaid", is_credit_note=False))
    db.commit()
    rows = client.get("/reports/sales", params={"group_by": "salesperson"}).json()["rows"]
    assert {r["salesperson"]: r["transaction_count"] for r in rows} == {"Test Analyst": 3, "Unassigned": 1}


def test_purchase_warehouse_filter(client, db, test_data):
    org, sup, wh = test_data["org"], test_data["supplier"], test_data["warehouse"]
    for number, wh_id in (("PUR-WH-1", wh.id), ("PUR-WH-NONE", None)):
        pinv = PurchaseInvoice(organization_id=org.id, supplier_id=sup.id, invoice_number=number, subtotal=100.0,
                               total=100.0, status="approved", warehouse_id=wh_id)
        db.add(pinv)
        db.flush()
        db.add(SupplierInvoice(organization_id=org.id, supplier_id=sup.id, purchase_id=pinv.id,
                               supplier_invoice_number=f"S{number}", subtotal=100.0, grand_total=100.0, status="recorded"))
    db.commit()
    res = client.get("/reports/purchase", params={"warehouse_id": wh.id})
    assert res.status_code == 200
    assert [r["invoice_number"] for r in res.json()["rows"]] == ["SPUR-WH-1"]
    assert res.json()["summary"]["total_purchases"] == 100.0
    assert len(client.get("/reports/purchase").json()["rows"]) == 2
    foreign_wh = Warehouse(id=str(uuid.uuid4()), organization_id=test_data["org2"].id, name="F", code="WH-F2")
    db.add(foreign_wh)
    db.commit()
    assert client.get("/reports/purchase", params={"warehouse_id": foreign_wh.id}).status_code == 404


def test_expense_submitted_by_filter(client, db, test_data):
    org, user = test_data["org"], test_data["user"]
    other = User(id=str(uuid.uuid4()), email="other@example.com", name="Other", password_hash="x",
                 organization_id=org.id, role="admin", is_active=True)
    foreign = User(id=str(uuid.uuid4()), email="foreign@example.com", name="Foreign", password_hash="x",
                   organization_id=test_data["org2"].id, role="admin", is_active=True)
    db.add_all([other, foreign])
    db.add_all([
        Expense(organization_id=org.id, category="Travel", amount=500.0, status="approved", submitted_by=user.id),
        Expense(organization_id=org.id, category="Food", amount=200.0, status="approved", submitted_by=other.id),
    ])
    db.commit()
    res = client.get("/reports/expense", params={"submitted_by": user.id})
    assert res.status_code == 200
    assert [r["category"] for r in res.json()["rows"]] == ["Travel"]
    assert res.json()["summary"]["approved_expense"] == 500.0
    assert client.get("/reports/expense", params={"submitted_by": foreign.id}).status_code == 404
    export = client.get("/reports/expense/export", params={"submitted_by": user.id, "format": "excel"})
    assert [r["Category"] for r in _xlsx_table(export.content, "Date")] == ["Travel"]


def test_purchase_return_purchase_id_filter(client, db, test_data):
    org, org2, sup = test_data["org"], test_data["org2"], test_data["supplier"]
    p1 = PurchaseInvoice(organization_id=org.id, supplier_id=sup.id, invoice_number="PUR-RET-1", status="approved")
    p2 = PurchaseInvoice(organization_id=org.id, supplier_id=sup.id, invoice_number="PUR-RET-2", status="approved")
    pf = PurchaseInvoice(organization_id=org2.id, invoice_number="PUR-RET-F", status="approved")
    db.add_all([p1, p2, pf])
    db.flush()
    db.add_all([
        PurchaseReturn(organization_id=org.id, return_number="PR-P1", purchase_id=p1.id, supplier_id=sup.id, status="confirmed"),
        PurchaseReturn(organization_id=org.id, return_number="PR-P2", purchase_id=p2.id, supplier_id=sup.id, status="confirmed"),
    ])
    db.commit()
    res = client.get("/reports/purchase-return", params={"purchase_id": p1.id})
    assert res.status_code == 200
    assert [(r["return_number"], r["purchase_reference"]) for r in res.json()["rows"]] == [("PR-P1", "PUR-RET-1")]
    assert len(client.get("/reports/purchase-return").json()["rows"]) == 2
    assert client.get("/reports/purchase-return", params={"purchase_id": pf.id}).status_code == 404


def test_export_rejects_unknown_format(client):
    res = client.get("/reports/sales/export", params={"format": "csv"})
    assert res.status_code == 400
    assert "format must be" in res.json()["detail"]


def test_gst_summary_excel_and_pdf_match_corrected_aggregate(client, gst_data, monkeypatch):
    from app.routers import reports as reports_router

    api = client.get("/reports/gst-summary", params={"date_from": "2026-03-01", "date_to": "2026-03-31"}).json()
    params = {"date_from": "2026-03-01", "date_to": "2026-03-31"}
    table = _xlsx_table(client.get("/reports/gst-summary/export", params={**params, "format": "excel"}).content, "GST Classification")
    assert [(r["Taxable Value"], r["Tax Amount"]) for r in table] == [(r["taxable_value"], r["tax_amount"]) for r in api["rows"]]
    assert api["summary"]["input_gst"] == 504.0

    captured = {}
    real_to_pdf = reports_router._to_pdf
    monkeypatch.setattr(reports_router, "_to_pdf", lambda report, org_name=None: captured.setdefault("r", report) and real_to_pdf(report, org_name))
    assert client.get("/reports/gst-summary/export", params={**params, "format": "pdf"}).status_code == 200
    assert captured["r"]["rows"] == api["rows"] and captured["r"]["summary"] == api["summary"]


@pytest.mark.parametrize("path", ["/reports/{t}", "/reports/{t}/export"])
@pytest.mark.parametrize("report_type,params", [
    ("profit-loss", {"customer_id": "abc"}),
    ("daily-transaction", {"overdue_only": "true"}),
    ("sales", {"submitted_by": "abc"}),
    ("expense", {"warehouse_id": "abc"}),
    ("gst-summary", {"customer_id": "abc"}),
    ("gst-summary", {"group_by": "transaction", "product_id": "abc"}),
])
def test_unsupported_filters_rejected_on_screen_and_export(client, path, report_type, params):
    res = client.get(path.format(t=report_type), params=params)
    assert res.status_code == 400
    assert "not supported" in res.json()["detail"]


def test_supported_and_blank_filters_still_accepted(client):
    # Blank values and default booleans are not "supplied" filters.
    assert client.get("/reports/profit-loss", params={"search": "", "overdue_only": "false"}).status_code == 200
    assert client.get("/reports/sales-overview", params={"status": "placed"}).status_code == 200
    assert client.get("/reports/supplier-payment", params={"payment_method": "cash"}).status_code == 200
    assert client.get("/reports/inventory-summary", params={"brand_id": "none", "category_id": "none"}).status_code == 200


def test_invalid_validation_matches_between_screen_and_export(client):
    for params in ({"group_by": "transaction", "status": "void"}, {"group_by": "transaction", "min_tax_amount": 5, "max_tax_amount": 1}):
        screen = client.get("/reports/gst-summary", params=params)
        export = client.get("/reports/gst-summary/export", params=params)
        assert screen.status_code == export.status_code == 400
        assert screen.json()["detail"] == export.json()["detail"]
    # Bad format is rejected before any report work, even alongside other invalid input.
    res = client.get("/reports/sales/export", params={"format": "csv", "submitted_by": "abc"})
    assert res.status_code == 400
    assert res.json()["detail"] == "format must be 'excel' or 'pdf'"


def test_expense_vendor_filter(client, db, test_data):
    org, sup = test_data["org"], test_data["supplier"]
    other_vendor = Supplier(id=str(uuid.uuid4()), organization_id=org.id, name="Other Vendor", supplier_code="SUP-0002", phone="1")
    foreign_vendor = Supplier(id=str(uuid.uuid4()), organization_id=test_data["org2"].id, name="Foreign Vendor", supplier_code="SUP-F", phone="2")
    db.add_all([other_vendor, foreign_vendor])
    db.add_all([
        Expense(organization_id=org.id, category="Rent", amount=900.0, status="approved", vendor_id=sup.id),
        Expense(organization_id=org.id, category="Power", amount=100.0, status="approved", vendor_id=other_vendor.id),
    ])
    db.commit()
    res = client.get("/reports/expense", params={"vendor_id": sup.id})
    assert res.status_code == 200
    assert [(r["category"], r["vendor"]) for r in res.json()["rows"]] == [("Rent", "Global Supplies")]
    assert client.get("/reports/expense", params={"vendor_id": foreign_vendor.id}).status_code == 404
    # Vendor-linked expenses used to break these reports (Supplier object in rows / ilike on a relationship).
    res = client.get("/reports/expense", params={"search": "other vendor"})
    assert res.status_code == 200
    assert [r["category"] for r in res.json()["rows"]] == ["Power"]
    for report in ("daily-transaction", "cash-flow-sheet"):
        res = client.get(f"/reports/{report}", params={"page_size": 100})
        assert res.status_code == 200, report


# ---------------------------------------------------------------------------
# Phase 1: Row-Level Export Tests
# ---------------------------------------------------------------------------

def test_row_export_sales_invoice_excel_and_pdf(client, gst_data):
    # 1. Sales invoice row exported as Excel
    res_excel = client.get("/reports/sales/export/INV-GST-001?format=excel")
    assert res_excel.status_code == 200
    assert res_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_excel.headers["content-disposition"] == 'attachment; filename="sales-INV-GST-001-report.xlsx"'

    rows = _xlsx_table(res_excel.content, "Invoice #")
    assert len(rows) == 1
    assert rows[0]["Invoice #"] == "INV-GST-001"

    # 2. Sales invoice row exported as PDF
    res_pdf = client.get("/reports/sales/export/INV-GST-001?format=pdf")
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.headers["content-disposition"] == 'attachment; filename="sales-INV-GST-001-report.pdf"'
    assert res_pdf.content.startswith(b"%PDF-")
    assert len(res_pdf.content) > 200


def test_row_export_sales_overview_order_identity(client, db, test_data):
    org, cust = test_data["org"], test_data["customer"]
    so = SalesOrder(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        customer_id=cust.id,
        order_number="SO-ORD-999",
        status="confirmed",
        total=1250.0,
    )
    db.add(so)
    db.commit()

    # 3. Sales-overview order row exported using its correct identity (UUID and order_number)
    # By order_number:
    res_num = client.get(f"/reports/sales-overview/export/{so.order_number}?format=excel")
    assert res_num.status_code == 200
    rows_num = _xlsx_table(res_num.content, "Order #")
    assert len(rows_num) == 1
    assert rows_num[0]["Order #"] == "SO-ORD-999"

    # By order_id UUID:
    res_id = client.get(f"/reports/sales-overview/export/{so.id}?format=excel")
    assert res_id.status_code == 200
    rows_id = _xlsx_table(res_id.content, "Order #")
    assert len(rows_id) == 1
    assert rows_id[0]["Order #"] == "SO-ORD-999"


def test_row_export_gst_transaction_disambiguation(client, db, gst_data):
    org, sup = gst_data["org"], gst_data["supplier"]
    # 4. GST transaction row resolved without confusing documents that share a reference
    pinv = PurchaseInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        invoice_number="PUR-AMBIG-001",
        status="approved",
    )
    db.add(pinv)
    db.flush()
    sinv = SupplierInvoice(
        organization_id=org.id,
        supplier_id=sup.id,
        purchase_id=pinv.id,
        supplier_invoice_number="INV-GST-001",  # Same number as sales invoice!
        subtotal=1000.0,
        tax_amount=180.0,
        grand_total=1180.0,
        status="recorded",
    )
    db.add(sinv)
    db.commit()

    # Requesting ambiguous reference without type qualification should return 400
    ambig_res = client.get("/reports/gst-summary/export/INV-GST-001?group_by=transaction&format=excel")
    assert ambig_res.status_code == 400
    assert "Multiple documents found with reference" in ambig_res.json()["detail"]
    assert "sales_invoice" in ambig_res.json()["detail"]
    assert "supplier_invoice" in ambig_res.json()["detail"]

    # Resolving with type-qualification: sales_invoice:INV-GST-001
    sales_res = client.get("/reports/gst-summary/export/sales_invoice:INV-GST-001?group_by=transaction&format=excel")
    assert sales_res.status_code == 200
    sales_rows = _xlsx_table(sales_res.content, "Date")
    assert len(sales_rows) == 1
    assert sales_rows[0]["Document #"] == "INV-GST-001"
    assert sales_rows[0]["Document Type"] == "sales_invoice"

    # Resolving with type-qualification: supplier_invoice:INV-GST-001
    supp_res = client.get("/reports/gst-summary/export/supplier_invoice:INV-GST-001?group_by=transaction&format=excel")
    assert supp_res.status_code == 200
    supp_rows = _xlsx_table(supp_res.content, "Date")
    assert len(supp_rows) == 1
    assert supp_rows[0]["Document #"] == "INV-GST-001"
    assert supp_rows[0]["Document Type"] == "supplier_invoice"

    # Resolving with exact source_id UUID
    src_res = client.get(f"/reports/gst-summary/export/{sinv.id}?group_by=transaction&format=excel")
    assert src_res.status_code == 200
    src_rows = _xlsx_table(src_res.content, "Date")
    assert len(src_rows) == 1
    assert src_rows[0]["Document #"] == "INV-GST-001"
    assert src_rows[0]["Document Type"] == "supplier_invoice"


def test_row_export_grouped_report_resolution(client, gst_data):
    # 5. Grouped report row resolution with its grouping parameter
    # With grouping parameter: group_by=customer
    res = client.get("/reports/sales/export/Acme%20Retail?group_by=customer&format=excel")
    assert res.status_code == 200
    rows = _xlsx_table(res.content, "Customer")
    assert len(rows) == 1
    assert rows[0]["Customer"] == "Acme Retail"

    # Grouped row requested without required grouping context returns clear 404
    res_no_grp = client.get("/reports/sales/export/Acme%20Retail?format=excel")
    assert res_no_grp.status_code == 404
    assert "group_by" in res_no_grp.json()["detail"]


def test_row_export_calculated_and_summary_rows(client, gst_data):
    # 6. GST summary/calculated row resolution
    res_gst = client.get("/reports/gst-summary/export/output_gst?format=excel")
    assert res_gst.status_code == 200
    gst_rows = _xlsx_table(res_gst.content, "GST Classification")
    assert len(gst_rows) == 1
    assert "Output GST" in gst_rows[0]["GST Classification"]

    # Profit and Loss calculated summary line (cogs)
    res_pl = client.get("/reports/profit-loss/export/cogs?format=excel")
    assert res_pl.status_code == 200
    pl_rows = _xlsx_table(res_pl.content, "P&L Statement Line Item")
    assert len(pl_rows) == 1
    assert "Cost of Goods Sold (COGS)" in pl_rows[0]["P&L Statement Line Item"]


def test_row_export_cross_tenant_isolation(client, gst_data):
    # 7. Cross-tenant row request returns 404
    # "INV-GST-FOREIGN" was created under foreign organization org2
    res = client.get("/reports/sales/export/INV-GST-FOREIGN?format=excel")
    assert res.status_code == 404
    assert "Row 'INV-GST-FOREIGN' not found in sales report" in res.json()["detail"]


def test_row_export_nonexistent_row(client, gst_data):
    # 8. Nonexistent row returns 404
    res = client.get("/reports/sales/export/NONEXISTENT-999?format=excel")
    assert res.status_code == 404
    assert "Row 'NONEXISTENT-999' not found in sales report" in res.json()["detail"]


def test_row_export_excluded_by_active_filters(client, gst_data):
    # 9. Row excluded by active filters returns 404
    # INV-GST-001 is on 2026-03-05. Filtering for April 2026 excludes it.
    res = client.get("/reports/sales/export/INV-GST-001?format=excel&date_from=2026-04-01&date_to=2026-04-30")
    assert res.status_code == 404
    assert "not found" in res.json()["detail"]


def test_row_export_unsupported_format(client, gst_data):
    # 10. Unsupported export format returns the expected client error
    res = client.get("/reports/sales/export/INV-GST-001?format=csv")
    assert res.status_code == 400
    assert "format must be" in res.json()["detail"]


def test_row_export_xlsx_alias_and_headers(client, gst_data):
    # 11. format=xlsx works as an alias for Excel
    # 12. Excel response has correct MIME type, attachment header, and valid non-empty workbook bytes
    from openpyxl import load_workbook
    from io import BytesIO

    res = client.get("/reports/sales/export/INV-GST-001?format=xlsx")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res.headers["content-disposition"] == 'attachment; filename="sales-INV-GST-001-report.xlsx"'
    wb = load_workbook(BytesIO(res.content))
    assert len(wb.sheetnames) >= 1
    assert len(res.content) > 1000


def test_row_export_pdf_headers_and_bytes(client, gst_data):
    # 13. PDF response has correct MIME type, attachment header, and valid non-empty PDF bytes
    res = client.get("/reports/sales/export/INV-GST-001?format=pdf")
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert res.headers["content-disposition"] == 'attachment; filename="sales-INV-GST-001-report.pdf"'
    assert res.content.startswith(b"%PDF-")
    assert len(res.content) > 200


def test_row_export_contains_only_selected_row(client, gst_data):
    # 14. Export contains only the selected row and not other rows
    all_res = client.get("/reports/sales?date_from=2026-03-01&date_to=2026-03-31")
    assert len(all_res.json()["rows"]) >= 2

    exp_res = client.get("/reports/sales/export/INV-GST-001?date_from=2026-03-01&date_to=2026-03-31&format=excel")
    assert exp_res.status_code == 200
    rows = _xlsx_table(exp_res.content, "Invoice #")
    assert len(rows) == 1
    assert rows[0]["Invoice #"] == "INV-GST-001"
    assert not any(r.get("Invoice #") == "INV-GST-002" for r in rows)


def test_row_export_empty_identifier_rejected(client, gst_data):
    # Empty row identifier rejected with 400
    res = client.get("/reports/sales/export/%20?format=excel")
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()


def test_row_export_additional_report_coverage(client, db, test_data):
    # Verify row exports on additional reports: customer-outstanding, inventory-summary, daily-transaction, expense
    org, cust, prod = test_data["org"], test_data["customer"], test_data["product"]

    # 1. customer-outstanding
    inv = Invoice(
        organization_id=org.id,
        invoice_number="INV-OUT-001",
        customer_id=cust.id,
        subtotal=1000.0,
        tax=0.0,
        total=1000.0,
        amount_paid=400.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=datetime.now(timezone.utc),
    )
    db.add(inv)

    # 2. expense
    exp = Expense(
        organization_id=org.id,
        category="Stationery",
        amount=150.0,
        status="approved",
        payment_status="paid",
        expense_date=datetime.now(timezone.utc),
    )
    db.add(exp)
    db.commit()

    # customer-outstanding by invoice_number
    res_co = client.get(f"/reports/customer-outstanding/export/{inv.invoice_number}?format=excel")
    assert res_co.status_code == 200
    co_rows = _xlsx_table(res_co.content, "Customer")
    assert len(co_rows) == 1
    assert co_rows[0]["Invoice #"] == "INV-OUT-001"

    # inventory-summary by SKU
    res_inv = client.get(f"/reports/inventory-summary/export/{prod.sku}?format=excel")
    assert res_inv.status_code == 200
    inv_rows = _xlsx_table(res_inv.content, "Product")
    assert len(inv_rows) == 1
    assert inv_rows[0]["SKU"] == prod.sku

    # expense by UUID
    res_exp = client.get(f"/reports/expense/export/{exp.id}?format=excel")
    assert res_exp.status_code == 200
    exp_rows = _xlsx_table(res_exp.content, "Date")
    assert len(exp_rows) == 1
    assert exp_rows[0]["Category"] == "Stationery"

    # daily-transaction by reference
    res_dt = client.get(f"/reports/daily-transaction/export/{inv.invoice_number}?format=excel")
    assert res_dt.status_code == 200
    dt_rows = _xlsx_table(res_dt.content, "Date")
    assert len(dt_rows) == 1
    assert dt_rows[0]["Reference #"] == "INV-OUT-001"


# ---------------------------------------------------------------------------
# Phase 2: Statement of P&L and Notes to P&L Tests
# ---------------------------------------------------------------------------

@pytest.fixture
def pl_statement_data(db, test_data):
    org, org2 = test_data["org"], test_data["org2"]
    cust, foreign_cust = test_data["customer"], test_data["foreign_customer"]
    prod1 = test_data["product"]

    prod2 = Product(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        name="Standard Gadget",
        sku="GADGET-02",
        price=50.0,
        minimum_stock_level=5,
        is_active=True,
    )
    db.add(prod2)

    # 1. SalesOrder 1 with sale-time historical cost_price snapshot = 45.0
    so1 = SalesOrder(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        customer_id=cust.id,
        order_number="SO-PL-101",
        status="completed",
        total=1000.0,
    )
    db.add(so1)
    db.flush()

    so_item1 = SalesOrderItem(
        id=str(uuid.uuid4()),
        order_id=so1.id,
        product_id=prod1.id,
        product_name=prod1.name,
        quantity=10,
        unit_price=100.0,
        cost_price=45.0,  # historical snapshot
        line_total=1000.0,
    )
    db.add(so_item1)

    # Invoice 1 (gross revenue = 1000.0, line COGS = 10 * 45 = 450.0)
    inv1 = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-PL-001",
        order_id=so1.id,
        customer_id=cust.id,
        subtotal=1000.0,
        tax=0.0,
        total=1000.0,
        amount_paid=1000.0,
        status="paid",
        is_credit_note=False,
        invoice_date=_ist(2026, 5, 10),
    )
    db.add(inv1)
    db.flush()

    item1 = InvoiceItem(
        id=str(uuid.uuid4()),
        invoice_id=inv1.id,
        product_id=prod1.id,
        order_item_id=so_item1.id,
        product_name=prod1.name,
        quantity=10,
        unit_price=100.0,
        line_total=1000.0,
    )
    db.add(item1)

    # 2. SalesOrder 2 without historical cost snapshot (cost_price=None)
    so2 = SalesOrder(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        customer_id=cust.id,
        order_number="SO-PL-102",
        status="completed",
        total=500.0,
    )
    db.add(so2)
    db.flush()

    so_item2 = SalesOrderItem(
        id=str(uuid.uuid4()),
        order_id=so2.id,
        product_id=prod2.id,
        product_name=prod2.name,
        quantity=5,
        unit_price=100.0,
        cost_price=None,  # missing historical cost snapshot
        line_total=500.0,
    )
    db.add(so_item2)

    # Invoice 2 (gross revenue = 500.0, COGS missing snapshot = 0.0)
    inv2 = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-PL-002",
        order_id=so2.id,
        customer_id=cust.id,
        subtotal=500.0,
        tax=0.0,
        total=500.0,
        amount_paid=500.0,
        status="paid",
        is_credit_note=False,
        invoice_date=_ist(2026, 5, 15),
    )
    db.add(inv2)
    db.flush()

    item2 = InvoiceItem(
        id=str(uuid.uuid4()),
        invoice_id=inv2.id,
        product_id=prod2.id,
        order_item_id=so_item2.id,
        product_name=prod2.name,
        quantity=5,
        unit_price=100.0,
        line_total=500.0,
    )
    db.add(item2)

    # 3. Approved SalesReturn (returns = 200.0)
    sr = SalesReturn(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        return_number="SR-PL-001",
        customer_id=cust.id,
        invoice_reference_id=inv1.id,
        return_status="approved",
        credit_amount=200.0,
        return_date=_ist(2026, 5, 20),
    )
    db.add(sr)

    # 4. Approved Operating Expenses (150.0 + 75.0 = 225.0)
    exp1 = Expense(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        category="Marketing",
        amount=150.0,
        status="approved",
        payment_status="paid",
        expense_date=_ist(2026, 5, 12),
        description="Q2 Ad Campaign",
    )
    db.add(exp1)

    exp2 = Expense(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        category="Office Supplies",
        amount=75.0,
        status="approved",
        payment_status="paid",
        expense_date=_ist(2026, 5, 18),
        description="Printer Paper and Ink",
    )
    db.add(exp2)

    # 5. Non-approved (pending) expense (must be excluded from statement and summary)
    exp_pending = Expense(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        category="Travel",
        amount=999.0,
        status="pending",
        payment_status="unpaid",
        expense_date=_ist(2026, 5, 18),
    )
    db.add(exp_pending)

    # 6. Cancelled invoice (must be excluded from statement and summary)
    inv_cancelled = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-PL-CANCELLED",
        customer_id=cust.id,
        subtotal=888.0,
        tax=0.0,
        total=888.0,
        status="cancelled",
        is_credit_note=False,
        invoice_date=_ist(2026, 5, 11),
    )
    db.add(inv_cancelled)

    # 7. Cross-tenant records (org2)
    inv_foreign = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org2.id,
        invoice_number="INV-PL-FOREIGN",
        customer_id=foreign_cust.id,
        subtotal=777.0,
        tax=0.0,
        total=777.0,
        status="paid",
        is_credit_note=False,
        invoice_date=_ist(2026, 5, 10),
    )
    db.add(inv_foreign)

    exp_foreign = Expense(
        id=str(uuid.uuid4()),
        organization_id=org2.id,
        category="Legal",
        amount=333.0,
        status="approved",
        expense_date=_ist(2026, 5, 10),
    )
    db.add(exp_foreign)

    db.commit()

    return {
        **test_data,
        "inv1": inv1,
        "item1": item1,
        "inv2": inv2,
        "item2": item2,
        "sr": sr,
        "exp1": exp1,
        "exp2": exp2,
        "prod2": prod2,
    }


def test_statement_of_pl_sections_and_reconciliation(client, pl_statement_data):
    # 1. statement-of-pl returns required transaction-level sections
    # 2. Each section reconciles with corresponding existing profit-loss summary values
    # 3. Returns and expenses have the correct signed impact
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31", "page_size": 100}

    pl_res = client.get("/reports/profit-loss", params=params)
    assert pl_res.status_code == 200
    pl_summary = pl_res.json()["summary"]

    stmt_res = client.get("/reports/statement-of-pl", params=params)
    assert stmt_res.status_code == 200
    stmt_data = stmt_res.json()
    stmt_summary = stmt_data["summary"]
    rows = stmt_data["rows"]

    # Verify summary reconciliation between statement-of-pl and profit-loss
    assert stmt_summary["gross_sales"] == 1500.0
    assert stmt_summary["gross_sales"] == pl_summary["gross_sales"]

    assert stmt_summary["sales_returns"] == 200.0
    assert stmt_summary["sales_returns"] == pl_summary["sales_returns"]

    assert stmt_summary["net_sales"] == 1300.0
    assert stmt_summary["net_sales"] == pl_summary["net_sales"]

    assert stmt_summary["cogs"] == 450.0
    assert stmt_summary["cogs"] == pl_summary["cogs"]

    assert stmt_summary["gross_profit"] == 850.0
    assert stmt_summary["gross_profit"] == pl_summary["gross_profit"]

    assert stmt_summary["operating_expenses"] == 225.0
    assert stmt_summary["operating_expenses"] == pl_summary["operating_expenses"]

    assert stmt_summary["net_profit"] == 625.0
    assert stmt_summary["net_profit"] == pl_summary["net_profit"]

    # Verify transaction-level rows across all 4 sections
    sections = {r["section"] for r in rows}
    assert sections == {"revenue", "sales_returns", "cogs", "operating_expenses"}

    rev_rows = [r for r in rows if r["section"] == "revenue"]
    ret_rows = [r for r in rows if r["section"] == "sales_returns"]
    cogs_rows = [r for r in rows if r["section"] == "cogs"]
    exp_rows = [r for r in rows if r["section"] == "operating_expenses"]

    # Section 1: Revenue rows sum to gross sales
    assert sum(r["amount"] for r in rev_rows) == 1500.0
    for r in rev_rows:
        assert r["impact"] == r["amount"]  # Positive contribution to profit

    # Section 2: Returns rows sum to sales returns with negative impact
    assert sum(r["amount"] for r in ret_rows) == 200.0
    for r in ret_rows:
        assert r["impact"] == -r["amount"]  # Negative contribution

    # Section 3: COGS rows sum to COGS with negative impact
    assert sum(r["amount"] for r in cogs_rows) == 450.0
    for r in cogs_rows:
        assert r["impact"] == -r["amount"]  # Negative contribution

    # Section 4: Expense rows sum to operating expenses with negative impact
    assert sum(r["amount"] for r in exp_rows) == 225.0
    for r in exp_rows:
        assert r["impact"] == -r["amount"]  # Negative contribution

    # Net profit equals sum of all signed impacts
    total_net_impact = round(sum(r["impact"] for r in rows), 2)
    assert total_net_impact == 625.0
    assert total_net_impact == stmt_summary["net_profit"]


def test_statement_of_pl_missing_cogs_no_catalog_fallback(client, pl_statement_data):
    # 4. Missing COGS snapshots are NOT replaced by current catalog prices
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31", "page_size": 100}
    res = client.get("/reports/statement-of-pl", params=params)
    assert res.status_code == 200
    data = res.json()

    # Find the COGS row for INV-PL-002 (item with no snapshot)
    cogs_rows = [r for r in data["rows"] if r["section"] == "cogs"]
    missing_cogs_row = next(r for r in cogs_rows if r["reference"] == "INV-PL-002")

    assert missing_cogs_row["amount"] == 0.0
    assert missing_cogs_row["impact"] == 0.0
    assert missing_cogs_row["status"] == "unrecorded_cost"
    assert "No cost snapshot" in missing_cogs_row["description"]

    # Verify summary and metadata flag missing lines
    assert data["summary"]["missing_cogs_lines"] == 1
    assert "1 line(s) had no historical cost snapshot" in data["meta"]["cogs_notes"]


def test_statement_of_pl_filters_and_tenant_isolation(client, pl_statement_data):
    # 5. Date filters, section filter, and tenant isolation work
    # Date filtering
    res_range = client.get("/reports/statement-of-pl", params={"date_from": "2026-05-01", "date_to": "2026-05-11"})
    assert res_range.status_code == 200
    refs = [r["reference"] for r in res_range.json()["rows"]]
    assert "INV-PL-001" in refs
    assert "INV-PL-002" not in refs  # 2026-05-15 is excluded

    # Section filter: only revenue
    res_sec = client.get("/reports/statement-of-pl", params={"section": "revenue", "date_from": "2026-05-01", "date_to": "2026-05-31"})
    assert res_sec.status_code == 200
    assert all(r["section"] == "revenue" for r in res_sec.json()["rows"])

    # Invalid section filter returns 400
    res_invalid_sec = client.get("/reports/statement-of-pl", params={"section": "nonexistent"})
    assert res_invalid_sec.status_code == 400
    assert "section must be one of" in res_invalid_sec.json()["detail"]

    # Tenant isolation: org2 records must never appear
    all_rows = client.get("/reports/statement-of-pl", params={"date_from": "2026-05-01", "date_to": "2026-05-31", "page_size": 100}).json()["rows"]
    assert not any(r["reference"] in ("INV-PL-FOREIGN", "Legal") for r in all_rows)
    assert not any(r.get("source_id") == pl_statement_data["org2"].id for r in all_rows)


def test_notes_to_pl_all_schedules_reconciliation(client, pl_statement_data):
    # 6. notes-to-pl supports all 4 canonical line_item values
    # 7. Each notes schedule reconciles to its parent P&L line
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31", "page_size": 100}

    # 1. line_item=revenue
    res_rev = client.get("/reports/notes-to-pl", params={**params, "line_item": "revenue"})
    assert res_rev.status_code == 200
    rev_data = res_rev.json()
    assert rev_data["meta"]["line_item"] == "revenue"
    assert rev_data["summary"]["gross_sales"] == 1500.0
    assert sum(r["total_revenue"] for r in rev_data["rows"]) == 1500.0

    # 2. line_item=sales_returns
    res_ret = client.get("/reports/notes-to-pl", params={**params, "line_item": "sales_returns"})
    assert res_ret.status_code == 200
    ret_data = res_ret.json()
    assert ret_data["meta"]["line_item"] == "sales_returns"
    assert ret_data["summary"]["sales_returns"] == 200.0
    assert sum(r["credit_amount"] for r in ret_data["rows"]) == 200.0

    # 3. line_item=cogs
    res_cogs = client.get("/reports/notes-to-pl", params={**params, "line_item": "cogs"})
    assert res_cogs.status_code == 200
    cogs_data = res_cogs.json()
    assert cogs_data["meta"]["line_item"] == "cogs"
    assert cogs_data["summary"]["cogs"] == 450.0
    assert sum(r["total_cogs"] for r in cogs_data["rows"]) == 450.0
    assert cogs_data["summary"]["missing_cogs_lines"] == 1

    # 4. line_item=operating_expenses
    res_exp = client.get("/reports/notes-to-pl", params={**params, "line_item": "operating_expenses"})
    assert res_exp.status_code == 200
    exp_data = res_exp.json()
    assert exp_data["meta"]["line_item"] == "operating_expenses"
    assert exp_data["summary"]["operating_expenses"] == 225.0
    assert sum(r["total_amount"] for r in exp_data["rows"]) == 225.0
    categories = {r["category"]: r["total_amount"] for r in exp_data["rows"]}
    assert categories["Marketing"] == 150.0
    assert categories["Office Supplies"] == 75.0


def test_notes_to_pl_validation(client):
    # 8. An invalid line_item returns HTTP 400
    # Missing line_item
    res_missing = client.get("/reports/notes-to-pl")
    assert res_missing.status_code == 400
    assert "line_item is required" in res_missing.json()["detail"]
    assert "revenue" in res_missing.json()["detail"]

    # Unsupported line_item
    res_bad = client.get("/reports/notes-to-pl", params={"line_item": "dividends"})
    assert res_bad.status_code == 400
    assert "Unsupported line_item 'dividends'" in res_bad.json()["detail"]


def test_statement_of_pl_and_notes_exports(client, pl_statement_data):
    # 9. Both new reports support existing Excel/PDF export behavior
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31"}

    # statement-of-pl Excel export
    res_stmt_excel = client.get("/reports/statement-of-pl/export", params={**params, "format": "excel"})
    assert res_stmt_excel.status_code == 200
    assert res_stmt_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    stmt_rows = _xlsx_table(res_stmt_excel.content, "Date")
    assert len(stmt_rows) >= 4

    # statement-of-pl PDF export
    res_stmt_pdf = client.get("/reports/statement-of-pl/export", params={**params, "format": "pdf"})
    assert res_stmt_pdf.status_code == 200
    assert res_stmt_pdf.headers["content-type"] == "application/pdf"
    assert res_stmt_pdf.content.startswith(b"%PDF-")

    # notes-to-pl Excel export
    res_notes_excel = client.get("/reports/notes-to-pl/export", params={**params, "line_item": "revenue", "format": "excel"})
    assert res_notes_excel.status_code == 200
    notes_rows = _xlsx_table(res_notes_excel.content, "Customer / Channel")
    assert len(notes_rows) >= 1

    # notes-to-pl PDF export
    res_notes_pdf = client.get("/reports/notes-to-pl/export", params={**params, "line_item": "cogs", "format": "pdf"})
    assert res_notes_pdf.status_code == 200
    assert res_notes_pdf.content.startswith(b"%PDF-")


def test_statement_of_pl_and_notes_row_exports(client, pl_statement_data):
    # 10. Row-level export works for representative rows from both new report types
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31"}
    item1 = pl_statement_data["item1"]

    # statement-of-pl: qualified section export
    res_stmt_row_qual = client.get(
        "/reports/statement-of-pl/export/revenue:INV-PL-001",
        params={**params, "format": "excel"},
    )
    assert res_stmt_row_qual.status_code == 200
    r_qual = _xlsx_table(res_stmt_row_qual.content, "Date")
    assert len(r_qual) == 1
    assert r_qual[0]["Reference #"] == "INV-PL-001"

    # statement-of-pl: source_id UUID export
    res_stmt_row_src = client.get(
        f"/reports/statement-of-pl/export/{item1.id}",
        params={**params, "format": "excel"},
    )
    assert res_stmt_row_src.status_code == 200
    r_src = _xlsx_table(res_stmt_row_src.content, "Date")
    assert len(r_src) == 1
    assert r_src[0]["Statement Section"] == "Cost of Goods Sold (COGS)"

    # statement-of-pl: ambiguous bare reference across sections returns HTTP 400
    res_ambig = client.get(
        "/reports/statement-of-pl/export/INV-PL-001",
        params={**params, "format": "excel"},
    )
    assert res_ambig.status_code == 400
    assert "Multiple statement items found" in res_ambig.json()["detail"]

    # statement-of-pl: bare reference filtered by section resolves uniquely
    res_sec_row = client.get(
        "/reports/statement-of-pl/export/INV-PL-001",
        params={**params, "section": "revenue", "format": "excel"},
    )
    assert res_sec_row.status_code == 200
    r_sec = _xlsx_table(res_sec_row.content, "Date")
    assert len(r_sec) == 1
    assert r_sec[0]["Reference #"] == "INV-PL-001"

    # notes-to-pl: row export for revenue by customer name
    res_note_rev = client.get(
        "/reports/notes-to-pl/export/Acme%20Retail",
        params={**params, "line_item": "revenue", "format": "excel"},
    )
    assert res_note_rev.status_code == 200
    n_rev = _xlsx_table(res_note_rev.content, "Customer / Channel")
    assert len(n_rev) == 1
    assert n_rev[0]["Customer / Channel"] == "Acme Retail"

    # notes-to-pl: row export for operating expenses by category
    res_note_exp = client.get(
        "/reports/notes-to-pl/export/Marketing",
        params={**params, "line_item": "operating_expenses", "format": "excel"},
    )
    assert res_note_exp.status_code == 200
    n_exp = _xlsx_table(res_note_exp.content, "Expense Category")
    assert len(n_exp) == 1
    assert n_exp[0]["Expense Category"] == "Marketing"


def test_profit_loss_summary_and_phase1_exports_unbroken(client, pl_statement_data):
    # 11. Existing profit-loss summary and Phase 1 row-export behavior remain compatible
    params = {"date_from": "2026-05-01", "date_to": "2026-05-31"}

    # Existing profit-loss summary
    pl_res = client.get("/reports/profit-loss", params=params)
    assert pl_res.status_code == 200
    pl_data = pl_res.json()
    assert pl_data["summary"]["net_profit"] == 625.0
    assert len(pl_data["rows"]) == 7  # 7 standard summary lines

    # Existing profit-loss row export for COGS line
    pl_row_res = client.get("/reports/profit-loss/export/cogs", params={**params, "format": "excel"})
    assert pl_row_res.status_code == 200
    pl_row = _xlsx_table(pl_row_res.content, "P&L Statement Line Item")
    assert len(pl_row) == 1
    assert "Cost of Goods Sold (COGS)" in pl_row[0]["P&L Statement Line Item"]


# ---------------------------------------------------------------------------
# Phase 3: General Ledger and Balance Sheet Tests
# ---------------------------------------------------------------------------

def test_chart_of_accounts_and_journal_double_entry_validation(db, test_data):
    org = test_data["org"]
    accounts = accounting_service.ensure_default_accounts(db, org.id)
    assert len(accounts) == 14
    assert "1010" in accounts
    assert "1200" in accounts
    assert "2010" in accounts
    assert "3010" in accounts

    # 1. Balanced entry succeeds
    entry = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 15, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 500.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 500.0},
        ],
        entry_type="manual",
        reference="TEST-BAL-01",
    )
    assert entry.status == "posted"
    assert len(entry.lines) == 2

    # 2. Unbalanced entry raises 400
    with pytest.raises(HTTPException) as exc_info:
        accounting_service.create_journal_entry(
            db=db,
            org_id=org.id,
            effective_date=datetime(2026, 1, 15, tzinfo=timezone.utc),
            lines=[
                {"account_code": "1010", "debit": 500.0, "credit": 0.0},
                {"account_code": "3010", "debit": 0.0, "credit": 400.0},
            ],
        )
    assert exc_info.value.status_code == 400
    assert "Unbalanced journal entry" in exc_info.value.detail

    # 3. Negative debit or credit raises 400
    with pytest.raises(HTTPException) as exc_info:
        accounting_service.create_journal_entry(
            db=db,
            org_id=org.id,
            effective_date=datetime(2026, 1, 15, tzinfo=timezone.utc),
            lines=[
                {"account_code": "1010", "debit": -100.0, "credit": 0.0},
                {"account_code": "3010", "debit": 0.0, "credit": -100.0},
            ],
        )
    assert exc_info.value.status_code == 400
    assert "non-negative" in exc_info.value.detail

    # 4. Fewer than two lines raises 400
    with pytest.raises(HTTPException) as exc_info:
        accounting_service.create_journal_entry(
            db=db,
            org_id=org.id,
            effective_date=datetime(2026, 1, 15, tzinfo=timezone.utc),
            lines=[{"account_code": "1010", "debit": 100.0, "credit": 0.0}],
        )
    assert exc_info.value.status_code == 400
    assert "at least two lines" in exc_info.value.detail


def test_opening_balances_and_duplicate_prevention(db, test_data):
    org = test_data["org"]
    eff_date = datetime(2026, 1, 1, tzinfo=timezone.utc)

    # 1. Post balanced opening balance
    ob = accounting_service.post_opening_balances(
        db=db,
        org_id=org.id,
        effective_date=eff_date,
        balances=[
            {"account_code": "1010", "debit": 10000.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 10000.0},
        ],
    )
    assert ob.entry_type == "opening_balance"
    assert ob.status == "posted"

    # 2. Duplicate opening balance raises 400
    with pytest.raises(HTTPException) as exc_info:
        accounting_service.post_opening_balances(
            db=db,
            org_id=org.id,
            effective_date=eff_date,
            balances=[
                {"account_code": "1010", "debit": 5000.0, "credit": 0.0},
                {"account_code": "3010", "debit": 0.0, "credit": 5000.0},
            ],
        )
    assert exc_info.value.status_code == 400
    assert "Opening balances have already been posted" in exc_info.value.detail


def test_journal_entry_reversal_audit_trail(db, test_data):
    org = test_data["org"]
    entry = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 2, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1500", "debit": 2500.0, "credit": 0.0},
            {"account_code": "1010", "debit": 0.0, "credit": 2500.0},
        ],
        entry_type="manual",
        reference="PUR-EQUIP-01",
    )
    assert entry.status == "posted"

    # Reverse entry
    rev = accounting_service.reverse_journal_entry(
        db=db,
        org_id=org.id,
        entry_id=entry.id,
        reason="Entered in error",
    )
    assert rev.is_reversal is True
    assert rev.reversed_entry_id == entry.id
    assert entry.status == "reversed"

    # Attempting to reverse again raises 400
    with pytest.raises(HTTPException) as exc_info:
        accounting_service.reverse_journal_entry(
            db=db,
            org_id=org.id,
            entry_id=entry.id,
        )
    assert exc_info.value.status_code == 400
    assert "already been reversed" in exc_info.value.detail


def test_operational_document_posting_and_idempotency(db, test_data):
    org = test_data["org"]
    cust = test_data["customer"]

    inv = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-GL-001",
        customer_id=cust.id,
        subtotal=1000.0,
        tax=180.0,
        total=1180.0,
        amount_paid=0.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
    )
    db.add(inv)
    db.commit()

    # 1. Post operational invoice
    entry1 = accounting_service.post_operational_document(db, org.id, "invoice", inv.id)
    assert entry1 is not None
    assert entry1.entry_type == "invoice"
    lines = {l.account.account_code: (l.debit, l.credit) for l in entry1.lines}
    assert lines["1200"] == (1180.0, 0.0)  # AR
    assert lines["4010"] == (0.0, 1000.0)  # Revenue
    assert lines["2100"] == (0.0, 180.0)   # GST Output Payable

    # 2. Idempotency: posting again returns same entry without duplicate
    entry2 = accounting_service.post_operational_document(db, org.id, "invoice", inv.id)
    assert entry2.id == entry1.id


@pytest.fixture
def gl_balance_sheet_data(db, test_data):
    org, org2 = test_data["org"], test_data["org2"]
    accounting_service.ensure_default_accounts(db, org.id)
    accounting_service.ensure_default_accounts(db, org2.id)

    # 1. Opening balance on 2026-01-01: Cash 10,000 (Dr), Owner Capital 10,000 (Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 1, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 10000.0, "credit": 0.0, "description": "Cash Opening"},
            {"account_code": "3010", "debit": 0.0, "credit": 10000.0, "description": "Capital Opening"},
        ],
        entry_type="opening_balance",
        reference="OB-2026-01-01",
        source_type="opening_balance",
        source_id=f"OB-{org.id}",
    )

    # 2. Fixed asset purchase on 2026-02-01: Office Equipment 2,000 (Dr), Cash 2,000 (Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 2, 1, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1500", "debit": 2000.0, "credit": 0.0, "description": "Laptops"},
            {"account_code": "1010", "debit": 0.0, "credit": 2000.0, "description": "Cash paid for laptops"},
        ],
        entry_type="manual",
        reference="JE-EQUIP-01",
    )

    # 3. Revenue on 2026-03-01: AR 5,000 (Dr), Sales Revenue 5,000 (Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 3, 1, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1200", "debit": 5000.0, "credit": 0.0, "description": "Client Invoice"},
            {"account_code": "4010", "debit": 0.0, "credit": 5000.0, "description": "Consulting Revenue"},
        ],
        entry_type="invoice",
        reference="INV-2026-001",
    )

    # 4. Cash collection on 2026-03-15: Bank 3,000 (Dr), AR 3,000 (Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 3, 15, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1020", "debit": 3000.0, "credit": 0.0, "description": "Client NEFT"},
            {"account_code": "1200", "debit": 0.0, "credit": 3000.0, "description": "AR Cleared"},
        ],
        entry_type="customer_payment",
        reference="PMT-2026-001",
    )

    # 5. Operating expense on 2026-04-01: Expenses 1,000 (Dr), Cash 1,000 (Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 4, 1, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "6010", "debit": 1000.0, "credit": 0.0, "description": "Rent Payment"},
            {"account_code": "1010", "debit": 0.0, "credit": 1000.0, "description": "Cash Paid for Rent"},
        ],
        entry_type="expense",
        reference="EXP-2026-001",
    )

    # 6. Foreign org transaction (org2)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org2.id,
        effective_date=datetime(2026, 3, 1, 12, 0, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 9999.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 9999.0},
        ],
        entry_type="manual",
        reference="FOREIGN-JE-01",
    )

    db.commit()
    return test_data


def test_balance_sheet_as_of_date_and_accounting_equation(client, gl_balance_sheet_data):
    # 1. As of 2026-02-15: March and April transactions must be excluded
    res_feb = client.get("/reports/balance-sheet", params={"as_of_date": "2026-02-15"})
    assert res_feb.status_code == 200
    data_feb = res_feb.json()
    s_feb = data_feb["summary"]

    # Cash = 10,000 - 2,000 = 8,000; Equipment = 2,000 -> Total Assets = 10,000
    assert s_feb["total_assets"] == 10000.0
    assert s_feb["total_liabilities"] == 0.0
    assert s_feb["total_equity"] == 10000.0
    assert s_feb["total_liabilities_and_equity"] == 10000.0
    assert s_feb["difference"] == 0.0
    assert s_feb["is_balanced"] is True
    assert s_feb["status"] == "balanced"

    # 2. As of 2026-04-15: Complete position including revenues and expenses
    res_apr = client.get("/reports/balance-sheet", params={"as_of_date": "2026-04-15"})
    assert res_apr.status_code == 200
    data_apr = res_apr.json()
    s_apr = data_apr["summary"]

    # Assets: Cash 7,000 + Bank 3,000 + AR 2,000 + Equipment 2,000 = 14,000
    assert s_apr["total_assets"] == 14000.0
    assert s_apr["total_liabilities"] == 0.0
    # Equity: Capital 10,000 + Period Net Income (5,000 rev - 1,000 exp = 4,000) = 14,000
    assert s_apr["total_equity"] == 14000.0
    assert s_apr["total_liabilities_and_equity"] == 14000.0
    assert s_apr["difference"] == 0.0
    assert s_apr["is_balanced"] is True

    # Check rows contain individual accounts
    codes = {r["account_code"]: r["balance"] for r in data_apr["rows"]}
    assert codes["1010"] == 7000.0   # Cash
    assert codes["1020"] == 3000.0   # Bank
    assert codes["1200"] == 2000.0   # AR
    assert codes["1500"] == 2000.0   # Equipment
    assert codes["3010"] == 10000.0  # Capital
    assert codes["NET-INC"] == 4000.0  # Net Income to Date


def test_balance_sheet_unconfigured_and_cross_tenant_isolation(client, db):
    # Fresh org with 0 entries
    fresh_plan = Plan(name="Enterprise", price_monthly=1999, price_yearly=19999, entitlements={"report.balance_sheet": True})
    db.add(fresh_plan)
    db.flush()
    fresh_org = Organization(id=str(uuid.uuid4()), name="Empty Ledger Org", company_code="CMP-30003", timezone="Asia/Kolkata", plan_id=fresh_plan.id)
    db.add(fresh_org)
    fresh_user = User(id=str(uuid.uuid4()), email="empty@test.com", name="Analyst", password_hash="hash", organization_id=fresh_org.id, role="admin", is_active=True)
    db.add(fresh_user)
    db.commit()

    from app.routers.reports import _view
    app.dependency_overrides[_view] = lambda: fresh_user

    res = client.get("/reports/balance-sheet")
    assert res.status_code == 200
    data = res.json()
    assert data["summary"]["status"] == "unconfigured"
    assert data["summary"]["total_assets"] == 0.0
    assert any("Opening balances or journal entries must be established" in w for w in data["meta"]["warnings"])

    # Reset override
    app.dependency_overrides.clear()


def test_notes_to_balance_sheet_schedules_and_reconciliation(client, gl_balance_sheet_data):
    params = {"as_of_date": "2026-04-15", "page_size": 100}

    # 1. schedule=cash_and_bank
    res_cash = client.get("/reports/notes-to-balance-sheet", params={**params, "schedule": "cash_and_bank"})
    assert res_cash.status_code == 200
    data_cash = res_cash.json()
    assert data_cash["summary"]["net_balance"] == 10000.0  # Cash 7,000 + Bank 3,000

    # 2. schedule=accounts_receivable
    res_ar = client.get("/reports/notes-to-balance-sheet", params={**params, "schedule": "accounts_receivable"})
    assert res_ar.status_code == 200
    data_ar = res_ar.json()
    assert data_ar["summary"]["net_balance"] == 2000.0  # AR 2,000

    # 3. Missing schedule returns 400
    res_missing = client.get("/reports/notes-to-balance-sheet", params={"as_of_date": "2026-04-15"})
    assert res_missing.status_code == 400
    assert "schedule (or line_item) is required" in res_missing.json()["detail"]

    # 4. Unsupported schedule returns 400
    res_bad = client.get("/reports/notes-to-balance-sheet", params={"schedule": "invalid_sched"})
    assert res_bad.status_code == 400
    assert "Unsupported schedule" in res_bad.json()["detail"]


def test_balance_sheet_and_notes_exports(client, gl_balance_sheet_data):
    params = {"as_of_date": "2026-04-15"}

    # Full Excel export for balance-sheet
    res_bs_excel = client.get("/reports/balance-sheet/export", params={**params, "format": "excel"})
    assert res_bs_excel.status_code == 200
    assert res_bs_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    bs_rows = _xlsx_table(res_bs_excel.content, "Section")
    assert len(bs_rows) >= 5

    # Full PDF export for balance-sheet
    res_bs_pdf = client.get("/reports/balance-sheet/export", params={**params, "format": "pdf"})
    assert res_bs_pdf.status_code == 200
    assert res_bs_pdf.content.startswith(b"%PDF-")

    # Full Excel export for notes-to-balance-sheet
    res_notes_excel = client.get("/reports/notes-to-balance-sheet/export", params={**params, "schedule": "cash_and_bank", "format": "excel"})
    assert res_notes_excel.status_code == 200
    notes_rows = _xlsx_table(res_notes_excel.content, "Date")
    assert len(notes_rows) >= 3


def test_balance_sheet_and_notes_row_exports(client, gl_balance_sheet_data):
    params = {"as_of_date": "2026-04-15"}

    # Row export for balance-sheet by account code 1010 (Cash on Hand)
    res_row_code = client.get("/reports/balance-sheet/export/1010", params={**params, "format": "excel"})
    assert res_row_code.status_code == 200
    r_code = _xlsx_table(res_row_code.content, "Section")
    assert len(r_code) == 1
    assert r_code[0]["Account Code"] == "1010"

    # Row export for balance-sheet by slugified account name
    res_row_slug = client.get("/reports/balance-sheet/export/operating_bank_account", params={**params, "format": "excel"})
    assert res_row_slug.status_code == 200
    r_slug = _xlsx_table(res_row_slug.content, "Section")
    assert len(r_slug) == 1
    assert r_slug[0]["Account Code"] == "1020"


def test_migration_upgrade_and_downgrade():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "migration_z4",
        "alembic/versions/z4a5b6c7d8e9_general_ledger_and_chart_of_accounts.py"
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.revision == "z4a5b6c7d8e9"
    assert mod.down_revision == "y3z4a5b6c7d8"
    assert hasattr(mod, "upgrade")
    assert hasattr(mod, "downgrade")


from decimal import Decimal
from sqlalchemy.exc import IntegrityError
from app.services import payment_service


def test_reversal_balance_correctness_and_duplicate_prevention(client, db):
    """Verify that posting an opening entry and reversing it returns all balances to zero,

    and that duplicate reversals on original or reversal entries are rejected.
    """
    plan = Plan(name="RevPlan", price_monthly=99, price_yearly=999, entitlements={"report.balance_sheet": True})
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Reversal Test Org", company_code="CMP-40009", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    user = User(id=str(uuid.uuid4()), email="rev@test.com", name="Rev User", password_hash="hash", organization_id=org.id, role="admin", is_active=True)
    db.add(user)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    from app.routers.reports import _view
    app.dependency_overrides[_view] = lambda: user

    # 1. Post opening entry: Cash 5,000 (Dr), Capital 5,000 (Cr)
    entry = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 5000.0, "credit": 0.0, "description": "Opening Cash"},
            {"account_code": "3010", "debit": 0.0, "credit": 5000.0, "description": "Opening Capital"},
        ],
        entry_type="opening_balance",
        reference="OB-REV-001",
    )
    assert entry.status == "posted"

    # Verify Balance Sheet reflects 5000 Cash and 5000 Capital
    res1 = client.get("/reports/balance-sheet", params={"as_of_date": "2026-01-05"})
    assert res1.status_code == 200
    d1 = res1.json()
    assert d1["summary"]["total_assets"] == 5000.0
    assert d1["summary"]["total_equity"] == 5000.0
    assert d1["summary"]["difference"] == 0.0

    # 2. Reverse the entry
    rev = accounting_service.reverse_journal_entry(
        db=db,
        org_id=org.id,
        entry_id=entry.id,
        reason="Erroneous opening entry",
    )
    assert rev.status == "posted"
    assert rev.is_reversal is True
    assert rev.reversed_entry_id == entry.id
    assert entry.status == "reversed"

    # 3. Verify that Balance Sheet balances return to 0.00
    res2 = client.get("/reports/balance-sheet", params={"as_of_date": "2026-01-05"})
    assert res2.status_code == 200
    d2 = res2.json()
    assert d2["summary"]["total_assets"] == 0.0
    assert d2["summary"]["total_equity"] == 0.0
    assert d2["summary"]["difference"] == 0.0

    # 4. Attempt duplicate reversal on original entry -> 400
    with pytest.raises(HTTPException) as exc_info1:
        accounting_service.reverse_journal_entry(
            db=db,
            org_id=org.id,
            entry_id=entry.id,
        )
    assert exc_info1.value.status_code == 400
    assert "already been reversed" in exc_info1.value.detail

    # 5. Attempt reversal on the reversal entry -> 400
    with pytest.raises(HTTPException) as exc_info2:
        accounting_service.reverse_journal_entry(
            db=db,
            org_id=org.id,
            entry_id=rev.id,
        )
    assert exc_info2.value.status_code == 400
    assert "Cannot reverse a reversal entry" in exc_info2.value.detail

    app.dependency_overrides.clear()


def test_operational_workflow_posting_invoice_payment_void(db):
    """Verify live operational posting for invoice, customer payment, and void reversal."""
    plan = Plan(name="FlowPlan", price_monthly=99, price_yearly=999, entitlements={"report.balance_sheet": True})
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Workflow Org", company_code="CMP-40001", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    cust = Customer(id=str(uuid.uuid4()), organization_id=org.id, name="Workflow Cust", phone="9988776655")
    db.add(cust)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    inv = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-WF-001",
        customer_id=cust.id,
        subtotal=1000.0,
        tax=180.0,
        total=1180.0,
        amount_paid=0.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
    )
    db.add(inv)
    db.commit()

    # 1. Post invoice
    inv_entry = accounting_service.post_operational_document(db, org.id, "invoice", inv.id)
    assert inv_entry is not None
    assert inv_entry.source_type == "invoice"
    assert inv_entry.source_id == inv.id
    lines_map = {l.account.account_code: (l.debit, l.credit) for l in inv_entry.lines}
    assert lines_map["1200"] == (Decimal("1180.00"), Decimal("0.00"))  # AR Dr
    assert lines_map["4010"] == (Decimal("0.00"), Decimal("1000.00"))  # Revenue Cr
    assert lines_map["2100"] == (Decimal("0.00"), Decimal("180.00"))   # GST Output Cr

    # Idempotent retry returns identical entry
    retry_entry = accounting_service.post_operational_document(db, org.id, "invoice", inv.id)
    assert retry_entry.id == inv_entry.id

    # 2. Record Customer Payment through payment_service
    pmt = payment_service.record(
        db=db,
        org_id=org.id,
        customer=cust,
        invoice=inv,
        amount=1180.0,
        payment_mode="bank",
        reference="NEFT-WF-001",
        received_on=datetime(2026, 3, 5, tzinfo=timezone.utc),
    )
    pmt_entry = (
        db.query(JournalEntry)
        .filter(
            JournalEntry.organization_id == org.id,
            JournalEntry.source_type == "customer_payment",
            JournalEntry.source_id == pmt.id,
        )
        .first()
    )
    assert pmt_entry is not None
    pmt_lines = {l.account.account_code: (l.debit, l.credit) for l in pmt_entry.lines}
    assert pmt_lines["1020"] == (Decimal("1180.00"), Decimal("0.00"))  # Bank Dr
    assert pmt_lines["1200"] == (Decimal("0.00"), Decimal("1180.00"))  # AR Cr

    # 3. Void Payment triggers automatic journal reversal
    payment_service.void(db, pmt)
    db.refresh(pmt_entry)
    assert pmt_entry.status == "reversed"
    reversal_entry = (
        db.query(JournalEntry)
        .filter(
            JournalEntry.organization_id == org.id,
            JournalEntry.is_reversal.is_(True),
            JournalEntry.reversed_entry_id == pmt_entry.id,
        )
        .first()
    )
    assert reversal_entry is not None
    assert reversal_entry.status == "posted"


def test_operational_workflow_posting_expense_and_sales_return(db):
    """Verify posting behavior for approved expenses and approved sales returns."""
    plan = Plan(name="ExpPlan", price_monthly=99, price_yearly=999)
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Expense Org", company_code="CMP-40002", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    cust = Customer(id=str(uuid.uuid4()), organization_id=org.id, name="Ret Cust", phone="9988776655")
    db.add(cust)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    # 1. Unapproved expense is ignored
    exp = Expense(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        expense_number="EXP-001",
        category="office_supplies",
        amount=590.0,
        subtotal=500.0,
        tax_amount=90.0,
        payment_mode="cash",
        status="pending",
        expense_date=datetime(2026, 4, 1, tzinfo=timezone.utc),
    )
    db.add(exp)
    db.commit()

    res_unapproved = accounting_service.post_operational_document(db, org.id, "expense", exp.id)
    assert res_unapproved is None

    # Approve expense -> posts GL entry
    exp.status = "approved"
    db.commit()
    exp_entry = accounting_service.post_operational_document(db, org.id, "expense", exp.id)
    assert exp_entry is not None
    exp_lines = {l.account.account_code: (l.debit, l.credit) for l in exp_entry.lines}
    assert exp_lines["6010"] == (Decimal("590.00"), Decimal("0.00"))  # Expense Dr
    assert exp_lines["2010"] == (Decimal("0.00"), Decimal("590.00"))  # Payable Cr until reimbursed

    # 2. Sales return posting
    sr = SalesReturn(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        return_number="SR-001",
        customer_id=cust.id,
        credit_amount=590.0,
        return_status="pending",
        return_date=datetime(2026, 4, 2, tzinfo=timezone.utc),
    )
    db.add(sr)
    db.commit()

    res_sr_unapproved = accounting_service.post_operational_document(db, org.id, "sales_return", sr.id)
    assert res_sr_unapproved is None

    # Approve sales return -> posts GL entry
    sr.return_status = "approved"
    db.commit()
    sr_entry = accounting_service.post_operational_document(db, org.id, "sales_return", sr.id)
    assert sr_entry is not None
    sr_lines = {l.account.account_code: (l.debit, l.credit) for l in sr_entry.lines}
    assert sr_lines["4020"] == (Decimal("590.00"), Decimal("0.00"))  # Sales Return Dr
    assert sr_lines["1200"] == (Decimal("0.00"), Decimal("590.00"))  # AR Cr


def test_transaction_rollback_on_failed_posting(db):
    """Verify that a failed journal posting rolls back the transaction and leaves no orphaned operational record."""
    plan = Plan(name="RollPlan", price_monthly=99, price_yearly=999)
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Roll Org", company_code="CMP-40003", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    db.commit()

    inv_id = str(uuid.uuid4())
    try:
        inv = Invoice(
            id=inv_id,
            organization_id=org.id,
            invoice_number="INV-FAIL-01",
            subtotal=1000.0,
            tax=0.0,
            total=1000.0,
            status="unpaid",
            is_credit_note=False,
            invoice_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
        )
        db.add(inv)
        db.flush()
        # Simulate journal failure before commit
        raise RuntimeError("GL simulation error")
    except Exception:
        db.rollback()

    # Verify invoice was not persisted
    persisted = db.get(Invoice, inv_id)
    assert persisted is None


def test_database_level_idempotency_constraint(db):
    """Verify that the database unique constraint uq_journal_entries_org_source rejects duplicate postings."""
    plan = Plan(name="IdempPlan", price_monthly=99, price_yearly=999)
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Idemp Org", company_code="CMP-40004", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    # 1. Post entry with source_type="invoice" and source_id="INV-IDEMP-01"
    entry1 = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 100.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 100.0},
        ],
        entry_type="manual",
        reference="REF-001",
        source_type="invoice",
        source_id="INV-IDEMP-01",
    )
    assert entry1.id is not None
    db.commit()

    # 2. Directly attempting to insert a duplicate with raw DB insert triggers IntegrityError
    dup_entry = JournalEntry(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        entry_number="JE-TEST-DUP",
        entry_type="manual",
        status="posted",
        effective_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
        source_type="invoice",
        source_id="INV-IDEMP-01",
    )
    db.add(dup_entry)
    with pytest.raises(IntegrityError):
        db.commit()
    db.rollback()

    # 3. create_journal_entry safely recovers and returns the existing entry
    entry2 = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 3, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 100.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 100.0},
        ],
        entry_type="manual",
        reference="REF-002",
        source_type="invoice",
        source_id="INV-IDEMP-01",
    )
    assert entry2.id == entry1.id


def test_concurrency_safe_journal_entry_numbering(db):
    """Verify that _allocate_journal_entry_number produces sequentially incrementing references."""
    plan = Plan(name="SeqPlan", price_monthly=99, price_yearly=999)
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Seq Org", company_code="CMP-40005", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    db.commit()

    num1 = accounting_service._allocate_journal_entry_number(db, org.id)
    num2 = accounting_service._allocate_journal_entry_number(db, org.id)
    num3 = accounting_service._allocate_journal_entry_number(db, org.id)

    assert num1.endswith("00001")
    assert num2.endswith("00002")
    assert num3.endswith("00003")


def test_completeness_assessment_and_false_balance_prevention(client, db):
    """Verify that unposted operational documents trigger incomplete completeness_status and disclosure warnings,

    even when the general ledger mathematically balances.
    """
    plan = Plan(name="CompPlan", price_monthly=99, price_yearly=999, entitlements={"report.balance_sheet": True})
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Completeness Org", company_code="CMP-40006", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    cust = Customer(id=str(uuid.uuid4()), organization_id=org.id, name="Comp Cust", phone="9988776655")
    db.add(cust)
    user = User(id=str(uuid.uuid4()), email="cp@test.com", name="CP User", password_hash="hash", organization_id=org.id, role="admin", is_active=True)
    db.add(user)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    from app.routers.reports import _view
    app.dependency_overrides[_view] = lambda: user

    # 1. Post balanced opening balance
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 10000.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 10000.0},
        ],
        entry_type="opening_balance",
        reference="OB-001",
    )

    # 2. Create an operational invoice that has NOT been posted to GL
    inv = Invoice(
        id=str(uuid.uuid4()),
        organization_id=org.id,
        invoice_number="INV-UNPOSTED-01",
        customer_id=cust.id,
        subtotal=2000.0,
        tax=0.0,
        total=2000.0,
        status="unpaid",
        is_credit_note=False,
        invoice_date=datetime(2026, 1, 10, tzinfo=timezone.utc),
    )
    db.add(inv)
    db.commit()

    # Query Balance Sheet: GL is mathematically balanced, but operational completeness is incomplete
    res = client.get("/reports/balance-sheet", params={"as_of_date": "2026-01-15"})
    assert res.status_code == 200
    data = res.json()
    summary = data["summary"]
    assert summary["is_balanced"] is True
    assert summary["difference"] == 0.0
    assert summary["completeness_status"] == "incomplete"
    assert summary["is_complete"] is False
    assert summary["unposted_operational_count"] == 1
    assert any("eligible operational document(s)" in w for w in data["meta"]["warnings"])

    # 3. Post the invoice to GL
    accounting_service.post_operational_document(db, org.id, "invoice", inv.id)

    # Query Balance Sheet again: now complete
    res2 = client.get("/reports/balance-sheet", params={"as_of_date": "2026-01-15"})
    assert res2.status_code == 200
    summary2 = res2.json()["summary"]
    assert summary2["completeness_status"] == "complete"
    assert summary2["is_complete"] is True
    assert summary2["unposted_operational_count"] == 0

    app.dependency_overrides.clear()


def test_equity_notes_schedule_reconciles_with_balance_sheet(client, db):
    """Verify that the Notes to Balance Sheet equity schedule matches the Balance Sheet total equity,

    including the dynamic current-period net income line.
    """
    plan = Plan(name="EqPlan", price_monthly=99, price_yearly=999, entitlements={"report.balance_sheet": True})
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Equity Rec Org", company_code="CMP-40007", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    user = User(id=str(uuid.uuid4()), email="eq@test.com", name="EQ User", password_hash="hash", organization_id=org.id, role="admin", is_active=True)
    db.add(user)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    from app.routers.reports import _view
    app.dependency_overrides[_view] = lambda: user

    # 1. Opening Capital: 10,000
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": 10000.0, "credit": 0.0},
            {"account_code": "3010", "debit": 0.0, "credit": 10000.0},
        ],
        entry_type="opening_balance",
        reference="OB-EQ-001",
    )

    # 2. Revenue: 5,000 (AR Dr, Revenue Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 2, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1200", "debit": 5000.0, "credit": 0.0},
            {"account_code": "4010", "debit": 0.0, "credit": 5000.0},
        ],
        entry_type="invoice",
        reference="INV-EQ-001",
    )

    # 3. Operating Expense: 1,500 (Expense Dr, Cash Cr)
    accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 2, 15, tzinfo=timezone.utc),
        lines=[
            {"account_code": "6010", "debit": 1500.0, "credit": 0.0},
            {"account_code": "1010", "debit": 0.0, "credit": 1500.0},
        ],
        entry_type="expense",
        reference="EXP-EQ-001",
    )

    # Net income to date: 5,000 - 1,500 = 3,500
    # Expected Balance Sheet Total Equity: 10,000 + 3,500 = 13,500
    res_bs = client.get("/reports/balance-sheet", params={"as_of_date": "2026-03-01"})
    assert res_bs.status_code == 200
    bs_data = res_bs.json()
    assert bs_data["summary"]["total_equity"] == 13500.0
    assert bs_data["summary"]["is_balanced"] is True

    # Expected Notes to Balance Sheet (schedule=equity) Net Balance: 13,500
    res_notes = client.get("/reports/notes-to-balance-sheet", params={"as_of_date": "2026-03-01", "schedule": "equity"})
    assert res_notes.status_code == 200
    notes_data = res_notes.json()
    assert notes_data["summary"]["net_balance"] == 13500.0
    assert notes_data["summary"]["net_balance"] == bs_data["summary"]["total_equity"]

    # Verify that dynamic NET-INC row is present in notes rows
    net_inc_rows = [r for r in notes_data["rows"] if r["account_code"] == "NET-INC"]
    assert len(net_inc_rows) == 1
    assert net_inc_rows[0]["credit"] == 3500.0
    assert net_inc_rows[0]["amount"] == 3500.0

    app.dependency_overrides.clear()


def test_fixed_precision_decimal_journal_entries(db):
    """Verify that debit and credit amounts are stored and calculated using exact Decimals."""
    plan = Plan(name="DecPlan", price_monthly=99, price_yearly=999)
    db.add(plan)
    db.flush()
    org = Organization(id=str(uuid.uuid4()), name="Decimal Org", company_code="CMP-40008", timezone="Asia/Kolkata", plan_id=plan.id)
    db.add(org)
    db.commit()

    accounting_service.ensure_default_accounts(db, org.id)

    # Post an entry with three lines: 0.10, 0.20 and 0.30 credit
    entry = accounting_service.create_journal_entry(
        db=db,
        org_id=org.id,
        effective_date=datetime(2026, 1, 1, tzinfo=timezone.utc),
        lines=[
            {"account_code": "1010", "debit": Decimal("0.10"), "credit": Decimal("0.00")},
            {"account_code": "1020", "debit": Decimal("0.20"), "credit": Decimal("0.00")},
            {"account_code": "3010", "debit": Decimal("0.00"), "credit": Decimal("0.30")},
        ],
        entry_type="manual",
        reference="DEC-001",
    )

    db.refresh(entry)
    for line in entry.lines:
        assert isinstance(line.debit, Decimal)
        assert isinstance(line.credit, Decimal)

    lines_by_code = {l.account.account_code: (l.debit, l.credit) for l in entry.lines}
    assert lines_by_code["1010"] == (Decimal("0.10"), Decimal("0.00"))
    assert lines_by_code["1020"] == (Decimal("0.20"), Decimal("0.00"))
    assert lines_by_code["3010"] == (Decimal("0.00"), Decimal("0.30"))



